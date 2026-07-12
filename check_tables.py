#!/usr/bin/env python3
"""全面检查每篇论文的表格，统计数量和序号，标记需要调整的。"""
import os, re, json
from docx import Document
from docx.oxml.ns import qn

MODIFIED_DIR = "/root/project/workspace/论文资料/整理/环艺_修改版"
files = sorted([f for f in os.listdir(MODIFIED_DIR) if f.endswith('.docx')])

results = []

for fname in files:
    path = os.path.join(MODIFIED_DIR, fname)
    doc = Document(path)
    
    tables_info = []
    
    # 遍历段落找到表格标题（表X-Y 或 表X： 或 "表X"等）
    # 同时记录每个表格在文档中的位置（段落序号）
    paras = doc.paragraphs
    
    # 方法1：通过正文中的"表X"标记查找表格标题
    caption_candidates = []
    for i, p in enumerate(paras):
        text = p.text.strip()
        if re.search(r'表\s*[\d]', text):
            # 可能是表格标题
            caption_candidates.append((i, text))
    
    # 方法2：直接分析每个表格，看表格前后的段落
    for t_idx, table in enumerate(doc.tables):
        # 标题检测
        title = ""
        title_before = ""
        title_after = ""
        
        # 找table在body中的位置
        tbl_element = table._tbl
        parent = tbl_element.getparent()
        body = tbl_element.getparent()
        
        # 找table的前一个兄弟元素
        prev_sib = tbl_element.getprevious()
        next_sib = tbl_element.getnext()
        
        if prev_sib is not None:
            # 检查前一个段落是否包含"表"
            prev_texts = []
            for p_elem in [prev_sib]:
                if p_elem.tag == qn('w:p'):
                    texts = [r.text or '' for r in p_elem.findall('.//' + qn('w:t'))]
                    prev_text = ''.join(texts).strip()
                    prev_texts.append(prev_text)
                    
                    if re.search(r'表\s*[\d]', prev_text):
                        title_before = prev_text
                    elif re.search(r'表\d', prev_text[:20]):
                        title_before = prev_text
        
        if next_sib is not None:
            if next_sib.tag == qn('w:p'):
                texts = [r.text or '' for r in next_sib.findall('.//' + qn('w:t'))]
                next_text = ''.join(texts).strip()
                if re.search(r'表\s*[\d]', next_text):
                    title_after = next_text
        
        # 行的数量
        n_rows = len(table.rows)
        n_cols = len(table.columns) if table.rows else 0
        
        # 第一行内容预览
        first_row = [cell.text.strip()[:20] for cell in table.rows[0].cells] if table.rows else []
        
        tables_info.append({
            'table_idx': t_idx,
            'rows': n_rows,
            'cols': n_cols,
            'title_before': title_before,
            'title_after': title_after,
            'first_row': first_row,
        })
    
    results.append({
        'file': fname,
        'n_tables': len(tables_info),
        'tables': tables_info,
    })
    
    print(f"\n{'='*60}")
    print(f"📄 {fname} — 共 {len(tables_info)} 个表格")
    print(f"{'='*60}")
    for t in tables_info:
        title_str = ""
        if t['title_before']: title_str = f"← 标题: {t['title_before'][:60]}"
        elif t['title_after']: title_str = f"→ 标题: {t['title_after'][:60]}"
        else: title_str = "⚠️ 无表格标题"
        
        print(f"  表格#{t['table_idx']+1}: {t['rows']}行×{t['cols']}列 {title_str}")
        print(f"    首行: {t['first_row'][:4]}")

print(f"\n\n{'='*60}")
print(f"📊 汇总：共检查 {len(results)} 篇论文")
tables_total = sum(r['n_tables'] for r in results)
print(f"📊 表格总数：{tables_total}")
