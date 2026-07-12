from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import textwrap
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]


def write_stub_modules(stub_root: Path) -> None:
    stubs = {
        "core.py": textwrap.dedent(
            """
            from __future__ import annotations

            import json
            import threading
            from pathlib import Path

            tasks_db = {}
            is_system_busy = False
            system_lock = threading.Lock()

            def finalize_docx(path: str):
                return path

            def save_diagnosis_report(data, path: str):
                Path(path).write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

            def extract_docx_text(path: str):
                return "stub docx text"

            def extract_cover_info(path: str):
                return {}
            """
        ),
        "profile.py": textwrap.dedent(
            """
            from __future__ import annotations

            def generate_profile(*args):
                return {"args": list(args)}
            """
        ),
        "generator.py": textwrap.dedent(
            """
            from __future__ import annotations

            def generate(profile, paper_type, update=None):
                return "stub paper text"
            """
        ),
        "renderer.py": textwrap.dedent(
            """
            from __future__ import annotations

            from pathlib import Path

            def render(input_txt, output_docx, paper_type, cover_info, drawing_folder=None, update=None):
                Path(output_docx).write_text("stub docx", encoding="utf-8")
                return output_docx
            """
        ),
        "reviser.py": textwrap.dedent(
            """
            from __future__ import annotations

            def _analyze_original_paper(original_text, paper_type):
                return {"original_text": original_text, "paper_type": paper_type}

            def _diagnose_and_reconstruct(analysis, paper_type):
                return {"analysis": analysis, "paper_type": paper_type}

            def _generate_revised(analysis, revision_plan, paper_type, update=None):
                return "stub revised paper text"
            """
        ),
        "batch_generate.py": textwrap.dedent(
            """
            from __future__ import annotations

            def batch_generate(input_json, output_dir, paper_type, limit, skip_existing):
                return {
                    "input_json": input_json,
                    "output_dir": output_dir,
                    "paper_type": paper_type,
                    "limit": limit,
                    "skip_existing": skip_existing,
                }
            """
        ),
    }

    for rel, content in stubs.items():
        (stub_root / rel).write_text(content.lstrip(), encoding="utf-8")


def build_env(stub_root: Path) -> dict[str, str]:
    env = os.environ.copy()
    pythonpath_parts = [str(stub_root), str(ROOT)]
    existing = env.get("PYTHONPATH")
    if existing:
        pythonpath_parts.append(existing)
    env["PYTHONPATH"] = os.pathsep.join(pythonpath_parts)
    env["PYTHONUNBUFFERED"] = "1"
    return env


def run_python(args: list[str], stub_root: Path, code: str | None = None) -> subprocess.CompletedProcess[str]:
    env = build_env(stub_root)
    cmd = [sys.executable, *args]
    if code is not None:
        cmd = [sys.executable, "-c", code]
    return subprocess.run(
        cmd,
        cwd=str(stub_root),
        env=env,
        text=True,
        capture_output=True,
    )


def run_python_real(args: list[str], code: str | None = None) -> subprocess.CompletedProcess[str]:
    env = os.environ.copy()
    env["PYTHONPATH"] = str(ROOT)
    env["PYTHONUNBUFFERED"] = "1"
    cmd = [sys.executable, *args]
    if code is not None:
        cmd = [sys.executable, "-c", code]
    return subprocess.run(
        cmd,
        cwd=str(ROOT),
        env=env,
        text=True,
        capture_output=True,
    )


def record(results: list[tuple[str, bool, str]], name: str, ok: bool, detail: str = "") -> None:
    results.append((name, ok, detail))


def check_help(stub_root: Path, module_name: str, expected: list[str] | None = None) -> tuple[bool, str]:
    proc = run_python(["-m", module_name, "--help"], stub_root)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    if expected:
        text = proc.stdout
        missing = [item for item in expected if item not in text]
        if missing:
            return False, f"missing={missing}; stdout={text.strip()}"
    return True, "ok"


def check_import_service(stub_root: Path) -> tuple[bool, str]:
    proc = run_python(
        [],
        stub_root,
        code="import delivery.service as s; print(s.app.title)",
    )
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    if "论文系统交付服务" not in proc.stdout:
        return False, f"unexpected stdout={proc.stdout.strip()}"
    return True, "ok"


def check_profile_help(stub_root: Path) -> tuple[bool, str]:
    proc = run_python_real(["/root/project/workspace/thesis-reviser-next/profile.py", "--help"])
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    text = proc.stdout
    if "土木" not in text:
        return False, f"missing civil option; stdout={text.strip()}"
    return True, "ok"


def check_batch_help(stub_root: Path) -> tuple[bool, str]:
    proc = run_python_real(["/root/project/workspace/thesis-reviser-next/batch_generate.py", "--help"])
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    text = proc.stdout
    if "土木" not in text:
        return False, f"missing civil option; stdout={text.strip()}"
    return True, "ok"


def check_renderer_help(stub_root: Path) -> tuple[bool, str]:
    proc = run_python_real(["/root/project/workspace/thesis-reviser-next/renderer.py", "--help"])
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    text = proc.stdout
    if "土木" not in text:
        return False, f"missing civil option; stdout={text.strip()}"
    return True, "ok"


def check_normalize_mapping(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        """
        import json
        from delivery.runner import normalize_paper_type

        cases = {
            "管理": "管理",
            "经管": "管理",
            "manage": "管理",
            "mg": "管理",
            "设计": "设计",
            "design": "设计",
            "sj": "设计",
            "机械": "机械",
            "mechanical": "机械",
            "mech": "机械",
            "mj": "机械",
            "土木": "土木",
            "civil": "土木",
            "cw": "土木",
            "": "管理",
            "未知": "管理",
        }

        result = {key: normalize_paper_type(key) for key in cases}
        print(json.dumps(result, ensure_ascii=False))
        for key, expected in cases.items():
            if result[key] != expected:
                raise SystemExit(f"{key!r} -> {result[key]!r}, expected {expected!r}")
        """
    ).strip()
    proc = run_python([], stub_root, code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_batch_parsing(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        """
        import json
        import sys
        from delivery import batch

        calls = {}

        def fake_standard(*args):
            calls["standard"] = list(args)

        def fake_civil(argv):
            calls["civil"] = list(argv or [])

        batch.run_standard_batch = fake_standard
        batch.run_civil_batch = fake_civil

        sys.argv = [
            "delivery.batch",
            "--input", "in.json",
            "--output-dir", "out",
            "--type", "设计",
            "--limit", "3",
            "--skip-existing",
        ]
        batch.main()
        print(json.dumps(calls, ensure_ascii=False))
        """
    ).strip()
    proc = run_python([], stub_root, code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    data = json.loads(proc.stdout.strip() or "{}")
    expected = ["in.json", "out", "设计", 3, True]
    if data.get("standard") != expected:
        return False, f"standard={data.get('standard')}; expected={expected}"

    civil_code = textwrap.dedent(
        """
        import json
        import sys
        from delivery import batch

        calls = {}

        def fake_standard(*args):
            calls["standard"] = list(args)

        def fake_civil(argv):
            calls["civil"] = list(argv or [])

        batch.run_standard_batch = fake_standard
        batch.run_civil_batch = fake_civil

        sys.argv = [
            "delivery.batch",
            "--workflow", "civil",
            "--civil-args",
            "--single", "张三",
        ]
        batch.main()
        print(json.dumps(calls, ensure_ascii=False))
        """
    ).strip()
    proc = run_python([], stub_root, code=civil_code)
    if proc.returncode != 0:
        return False, f"civil exit={proc.returncode}; stderr={proc.stderr.strip()}"
    data = json.loads(proc.stdout.strip() or "{}")
    if data.get("civil") != ["--single", "张三"]:
        return False, f"civil={data.get('civil')}; expected=['--single', '张三']"
    if "standard" in data:
        return False, f"unexpected standard call: {data['standard']}"
    return True, "ok"


def check_path_conventions(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        """
        import json
        import tempfile
        from pathlib import Path

        from delivery import runner, service

        with tempfile.TemporaryDirectory() as tmp:
            output_dir = Path(tmp) / "out"
            input_docx = Path(tmp) / "input.docx"
            input_docx.write_text("stub", encoding="utf-8")

            gen = runner.run_generate({"title": "t"}, "管理", output_dir)
            rev = runner.run_revise(str(input_docx), "管理", output_dir)
            folder = service._task_folder("task_demo")

            data = {
                "generate": {
                    "txt": Path(gen["txt_path"]).name,
                    "docx": Path(gen["docx_path"]).name,
                },
                "revise": {
                    "report": Path(rev["report_path"]).name,
                    "txt": Path(rev["txt_path"]).name,
                    "docx": Path(rev["docx_path"]).name,
                },
                "task_folder": str(folder),
            }
            print(json.dumps(data, ensure_ascii=False))
        """
    ).strip()
    proc = run_python([], stub_root, code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    data = json.loads(proc.stdout.strip() or "{}")
    if data.get("generate", {}).get("txt") != "00_完整论文.txt":
        return False, f"generate txt={data.get('generate', {}).get('txt')}"
    if data.get("generate", {}).get("docx") != "论文_含图表.docx":
        return False, f"generate docx={data.get('generate', {}).get('docx')}"
    if data.get("revise", {}).get("report") != "诊断报告.txt":
        return False, f"revise report={data.get('revise', {}).get('report')}"
    if data.get("revise", {}).get("txt") != "00_完整论文.txt":
        return False, f"revise txt={data.get('revise', {}).get('txt')}"
    if data.get("revise", {}).get("docx") != "论文_修改版.docx":
        return False, f"revise docx={data.get('revise', {}).get('docx')}"
    if not data.get("task_folder", "").endswith("/output/task_demo"):
        return False, f"task_folder={data.get('task_folder')}"
    return True, "ok"


def check_table_matrix_guard(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        """
        import json
        from core import canonicalize_design_table_payload

        normal = canonicalize_design_table_payload("项目,优化前,优化后", "成本控制,资金周转", "8|9;6|7")
        guard = canonicalize_design_table_payload("A,B", "x,y", "")

        print(json.dumps({"normal": normal, "guard": guard}, ensure_ascii=False))
        if guard["rows"] != [["x", ""], ["y", ""]]:
            raise SystemExit(f"guard rows unexpected: {guard['rows']!r}")
        """
    ).strip()
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    data = json.loads(proc.stdout.strip() or "{}")
    if data.get("normal", {}).get("rows") != [["成本控制", "8", "9"], ["资金周转", "6", "7"]]:
        return False, f"normal rows={data.get('normal', {}).get('rows')}"
    return True, "ok"


def check_multiline_table_parse(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        '''
        import json
        from core import tolerant_extract_tables

        text = """<table
        title='示例表'
        data='8|9;6|7'
        rows='成本控制,资金周转'
        header='项目,优化前,优化后'
        id='7'
        />"""
        tables = tolerant_extract_tables(text)
        print(json.dumps(tables, ensure_ascii=False))
        if not tables or tables[0]["title"] != "示例表":
            raise SystemExit(f"parse failed: {tables!r}")
        if tables[0]["rows"] != "成本控制,资金周转":
            raise SystemExit(f"rows failed: {tables[0]['rows']!r}")
        '''
    ).strip()
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_cover_aliases(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        """
        import json
        import tempfile
        from docx import Document
        from core import extract_cover_info

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/cover.docx"
            doc = Document()
            doc.add_paragraph("论文题目：A公司财务管理问题与对策研究")
            doc.add_paragraph("学生姓名：李四")
            doc.add_paragraph("准考证号：2026123456")
            doc.add_paragraph("指导老师：王老师")
            doc.add_paragraph("层次：专升本")
            doc.add_paragraph("2026年5月20日")
            doc.save(path)
            info = extract_cover_info(path)
            print(json.dumps(info, ensure_ascii=False))
            if info["student_id"] != "2026123456":
                raise SystemExit(f"student_id={info['student_id']!r}")
            if info["advisor"] != "王老师":
                raise SystemExit(f"advisor={info['advisor']!r}")
        """
    ).strip()
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_multiline_chart_parse(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        '''
        import json
        from core import tolerant_extract_charts

        text = """<chart
        data_source='测试数据'
        y='10,20,30'
        title='趋势图'
        x='2022,2023,2024'
        type='line'
        legend='销量'
        id='c1'
        unit='%'
        />"""
        charts = tolerant_extract_charts(text)
        print(json.dumps(charts, ensure_ascii=False))
        if not charts or charts[0]["title"] != "趋势图":
            raise SystemExit(f"parse failed: {charts!r}")
        if charts[0]["type"] != "line":
            raise SystemExit(f"type failed: {charts[0]['type']!r}")
        '''
    ).strip()
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_multiline_drawing_parse(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        """
        import json
        from core import extract_drawings_from_text

        text = '''<drawing
        id="d1"
        type="结构图"
        title="装配结构示意图"
        description="展示零件A与零件B的连接关系"
        />'''
        drawings = extract_drawings_from_text(text)
        print(json.dumps(drawings, ensure_ascii=False))
        if not drawings or drawings[0]["title"] != "装配结构示意图":
            raise SystemExit(f"parse failed: {drawings!r}")
        if drawings[0]["type"] != "结构图":
            raise SystemExit(f"type failed: {drawings[0]['type']!r}")
        """
    ).strip()
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_partial_drawing_parse(stub_root: Path) -> tuple[bool, str]:
    code = (
        "from core import extract_drawings_from_text\n"
        "text = '<drawing type=\"结构图\" description=\"展示零件A与零件B的连接关系\" />'\n"
        "drawings = extract_drawings_from_text(text)\n"
        "print(drawings)\n"
        "if not drawings:\n"
        "    raise SystemExit('no drawing parsed')\n"
        "if drawings[0]['title'] != '展示零件A与零件B的连接关系':\n"
        "    raise SystemExit(f\"unexpected title: {drawings[0]['title']!r}\")\n"
        "if drawings[0]['type'] != '结构图':\n"
        "    raise SystemExit(f\"unexpected type: {drawings[0]['type']!r}\")\n"
    )
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_placeholder_drawing_title_replace(stub_root: Path) -> tuple[bool, str]:
    code = (
        "from core import _canonicalize_drawing_line\n"
        "raw = '<drawing id=\"d2\" type=\"结构图\" title=\"示意图\" description=\"装配结构示意图\" />'\n"
        "parsed = _canonicalize_drawing_line(raw, 1, 0, len(raw))\n"
        "print(parsed)\n"
        "if not parsed:\n"
        "    raise SystemExit('parse failed')\n"
        "if parsed['title'] != '装配结构示意图':\n"
        "    raise SystemExit(f\"unexpected title: {parsed['title']!r}\")\n"
        "if parsed['type'] != '结构图':\n"
        "    raise SystemExit(f\"unexpected type: {parsed['type']!r}\")\n"
    )
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_mech_json_merge(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        """
        import json
        from core import merge_mech_json_into_drawings

        drawings = [
            {"id": "1", "seq": 1, "title": "原图", "description": "原描述", "type": "结构图"},
        ]
        mech_blocks = [
            {
                "payload": {
                    "machine_spec": {
                        "load": "500N",
                        "critical_params": [{"name": "载荷", "value": "500N"}],
                    },
                    "drawings": [
                        {
                            "id": "1",
                            "title": "结构图一",
                            "description": "结构说明",
                            "numeric_annotations": ["500N", "500N", "30mm"],
                        }
                    ],
                }
            }
        ]
        merged = merge_mech_json_into_drawings(drawings, mech_blocks)
        print(json.dumps(merged, ensure_ascii=False))
        ann = merged[0]["numeric_annotations"]
        if ann != ["load：500N", "载荷：500N", "500N", "30mm"]:
            raise SystemExit(f"annotations unexpected: {ann!r}")
        if merged[0]["title"] != "结构图一":
            raise SystemExit(f"title unexpected: {merged[0]['title']!r}")
        """
    ).strip()
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_manage_markup_fallback(stub_root: Path) -> tuple[bool, str]:
    code = (
        "from generator import normalize_manage_markup\n"
        "text = '''第1章 引言\n"
        "<chart id=\"1\" title=\"趋势图\" type=\"line\" x=\"A,B\" y=\"1,2\" />\n"
        "\n"
        "第2章 现状分析\n"
        "<chart id=\"2\" title=\"趋势图2\" type=\"line\" x=\"A,B\" y=\"3,4\" />\n"
        "'''\n"
        "out = normalize_manage_markup(text)\n"
        "print(out)\n"
        "if out.count(\"<table \") != 1:\n"
        "    raise SystemExit(f\"expected one fallback table, got {out.count('<table ')}\")\n"
        "if out.count(\"<chart \") != 2:\n"
        "    raise SystemExit(f\"expected two charts, got {out.count('<chart ')}\")\n"
    )
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_manage_markup_fallback_with_later_table(stub_root: Path) -> tuple[bool, str]:
    code = (
        "from generator import normalize_manage_markup\n"
        "text = '''第1章 引言\n"
        "<chart id=\"1\" title=\"趋势图\" type=\"line\" x=\"A,B\" y=\"1,2\" />\n"
        "\n"
        "第2章 现状分析\n"
        "<table id=\"9\" title=\"已有表\" header=\"A,B\" rows=\"x,y\" data=\"1|2;3|4\" />\n"
        "'''\n"
        "out = normalize_manage_markup(text)\n"
        "print(out)\n"
        "if out.count(\"<table \") != 2:\n"
        "    raise SystemExit(f\"expected two tables, got {out.count('<table ')}\")\n"
        "if out.count(\"<chart \") != 1:\n"
        "    raise SystemExit(f\"expected one chart, got {out.count('<chart ')}\")\n"
    )
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_table_tag_order_valid(stub_root: Path) -> tuple[bool, str]:
    code = (
        "from generator import _VALID_TABLE_TAG_RE\n"
        "valid = '<table title=\"T\" rows=\"R\" data=\"D\" id=\"1\" header=\"H\" />'\n"
        "invalid = '<table title=\"T\" rows=\"R\" data=\"D\" id=\"1\" />'\n"
        "print(bool(_VALID_TABLE_TAG_RE.match(valid)), bool(_VALID_TABLE_TAG_RE.match(invalid)))\n"
        "if not _VALID_TABLE_TAG_RE.match(valid):\n"
        "    raise SystemExit('valid tag should match')\n"
        "if _VALID_TABLE_TAG_RE.match(invalid):\n"
        "    raise SystemExit('invalid tag should not match')\n"
    )
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_annotation_wrapper_flatten(stub_root: Path) -> tuple[bool, str]:
    code = (
        "from core import _collect_annotation_lines\n"
        "payload = {'machine_spec': {'load': '500N', 'critical_params': [{'name': '载荷', 'value': '500N'}]}, 'drawings': [{'id': '1'}]}\n"
        "lines = _collect_annotation_lines(payload, 'payload')\n"
        "print(lines)\n"
        "if lines != ['load：500N', '载荷：500N']:\n"
        "    raise SystemExit(f'unexpected lines: {lines!r}')\n"
    )
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_civil_generation(stub_root: Path) -> tuple[bool, str]:
    code = (
        "import re\n"
        "import generator\n\n"
        "def fake_llm(prompt, max_tokens=0):\n"
        "    if '标签格式校验专家' in prompt:\n"
        "        return '{\"issues_found\": false, \"corrections\": []}'\n"
        "    if '参考文献' in prompt and '列出' in prompt:\n"
        "        return '[1] 张三. 土木工程设计研究[D]. 某出版社, 2024.'\n"
        "    if '摘要' in prompt and '关键词' in prompt:\n"
        "        return '本文围绕某市办公楼建筑与结构设计开展研究，完成工程概况、建筑设计、结构设计、结构计算与施工组织设计。\\n关键词：土木工程；框架结构；施工组织'\n"
        "    m = re.search(r'第(\\d+)章：([^\\n]+)', prompt)\n"
        "    if m:\n"
        "        chapter = m.group(2).strip()\n"
        "        return f'{chapter}\\n1.1 小节内容\\n<drawing id=\"1\" type=\"结构图\" title=\"图1-1 {chapter}示意图\" description=\"该图展示{chapter}相关参数，结构形式为框架结构，层高3.6m，建筑总高度22.8m。\"/>'\n"
        "    return '默认内容'\n\n"
        "old_llm = generator.call_llm\n"
        "generator.call_llm = fake_llm\n"
        "try:\n"
        "    txt = generator.generate({\n"
        "        'title': '某市办公楼建筑与结构设计',\n"
        "        'object_name': '某市办公楼',\n"
        "        'structure_type': '框架结构',\n"
        "        'major': '土木工程',\n"
        "    }, '土木')\n"
        "finally:\n"
        "    generator.call_llm = old_llm\n\n"
        "print(txt[:400])\n"
        "if '第1章 工程概况' not in txt:\n"
        "    raise SystemExit('missing chapter header')\n"
        "if '<drawing ' not in txt:\n"
        "    raise SystemExit('missing drawing tag')\n"
    )
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_civil_markup_normalization(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        """
        from generator import normalize_civil_markup

        raw = '''好的，遵照您的要求，我将以土木工程专业论文写作专家的身份，为您撰写毕业论文的第3章：结构设计。内容将严格遵循您提供的写作要点、结构参数和输出格式。
        第3章 结构设计
        <drawing id="1" type="结构图" title="图3-1 结构布置图" description="展示框架结构与柱网布置，层高3.6m，建筑总高度22.8m。"/>/
        <drawing />
        </draing
        /
        <table border="1" />
         <tr
         <th构件</th
         <th截面</th
         </tr
         <tr
         <td主梁</td
         <td300×600</td
         </tr
        </table
        正文段落内容。
        '''
        out = normalize_civil_markup(raw)
        print(out)
        bad = ['好的，', '作为土木工程专业论文写作专家', '遵照您的要求', '<drawing />', '</draing', '<draing', '<table border="1"']
        for token in bad:
            if token in out:
                raise SystemExit(f'found leaked token: {token!r}')
        if any(line.strip() == '/' for line in out.splitlines()):
            raise SystemExit('found standalone slash line')
        if '<drawing id="1"' not in out:
            raise SystemExit('missing canonical drawing tag')
        if '<table ' not in out:
            raise SystemExit('missing repaired table tag')
        """
    ).strip()
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_civil_label_refinement(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        """
        from generator import normalize_civil_markup

        raw = '''第2章 建筑设计
        <drawing id="1" type="drawing" title="标准层建筑平面图" description="展示二至四层标准层平面布局，包含阅览区、走廊、楼梯间、电梯厅、卫生间等功能分区及防火分区示意"/>
        第3章 结构设计
        <table id="1" title="表1" header="构件,截面尺寸 (mm),混凝土强度等级" rows="主梁,次梁" data="300×600|C30;200×400|C30"/>
        第5章 施工组织设计
        <drawing id="2" type="drawing" title="施工平面布置图" description="场地北侧为钢筋堆场与加工棚，南侧为模板堆场，东侧为1#塔吊，西侧为2#塔吊，西北角为办公区，西南角为生活区，环形施工道路贯穿全场"/>
        '''
        out = normalize_civil_markup(raw)
        print(out)
        if 'type="drawing"' in out:
            raise SystemExit('generic drawing type not refined')
        if 'type="施工平面布置图"' not in out:
            raise SystemExit('missing refined site plan type')
        if 'title="表3-1 主要结构构件截面及材料表"' not in out:
            raise SystemExit('missing refined structure table title')
        """
    ).strip()
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_civil_batch_and_render_support(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        '''
        import inspect
        import batch_generate
        import renderer

        entry = {"title": "某市办公楼建筑与结构设计", "company": "某市办公楼", "major": "土木工程", "context": "办公楼"}
        profile = batch_generate._build_profile(entry, "土木")
        print(profile)
        if profile.get("paper_type") != "土木":
            raise SystemExit(f"batch profile paper_type={profile.get('paper_type')!r}")

        src = inspect.getsource(renderer.render)
        if "土木" not in src:
            raise SystemExit("renderer missing civil branch")
        '''
    ).strip()
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_real_sample_structure_guard(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        """
        import re
        from collections import Counter
        from pathlib import Path

        root = Path('/root/project/workspace/thesis-reviser-next/output/real_four_type')
        samples = [
            ('管理', root / 'H市金桥物流有限公司仓储管理优化研究' / 'paper.txt'),
            ('设计', root / '基于老年人居家使用场景的折叠助行车产品优化设计研究' / 'paper.txt'),
            ('机械', root / '某立式加工中心自动换刀机械手设计与分析' / 'paper.txt'),
            ('土木', root / '某大学图书馆建筑与结构设计研究' / 'paper.txt'),
        ]

        expectations = {
            '管理': {
                3: {'tables': 2, 'charts': 2},
                4: {'tables': 2, 'charts': 1},
                5: {'tables': 4, 'charts': 1},
            },
            '设计': {
                3: {'tables': 1, 'drawings': 2, 'charts': 1},
                4: {'tables': 3, 'drawings': 5},
            },
            '机械': {
                1: {'drawings': 1},
                2: {'tables': 1, 'drawings': 1},
                3: {'drawings': 1},
                4: {'drawings': 6},
                5: {'tables': 1, 'drawings': 3},
                6: {'drawings': 5},
            },
            '土木': {
                2: {'drawings': 2},
                3: {'tables': 1},
                4: {'drawings': 4},
                5: {'tables': 2, 'drawings': 3},
            },
        }

        def chapter_sections(text: str):
            parts = text.split('---PAGE_BREAK---\\n')
            body = '\\n'.join(parts[2:]) if len(parts) >= 3 else text
            matches = list(re.finditer(r'^第(\\d+)章\\s*(.+)$', body, flags=re.M))
            out = []
            for i, m in enumerate(matches):
                start = m.end()
                end = matches[i + 1].start() if i + 1 < len(matches) else len(body)
                out.append((int(m.group(1)), m.group(2).strip(), body[start:end]))
            return out

        failures = []
        for paper_type, path in samples:
            text = path.read_text(encoding='utf-8')
            table_ids = re.findall(r'<table\\b[^>]*\\bid=\"([^\"]+)\"', text)
            drawing_ids = re.findall(r'<drawing\\b[^>]*\\bid=\"([^\"]+)\"', text)
            table_titles = re.findall(r'<table\\b[^>]*\\btitle=\"([^\"]+)\"', text)

            dup_table_ids = {k: v for k, v in Counter(table_ids).items() if v > 1}
            dup_drawing_ids = {k: v for k, v in Counter(drawing_ids).items() if v > 1}
            double_titles = [t for t in table_titles if re.match(r'^表\\d+\\s+表\\d+-\\d+', t)]

            if dup_table_ids:
                failures.append(f'{path.parent.name}: 重复表id {dup_table_ids}')
            if dup_drawing_ids:
                failures.append(f'{path.parent.name}: 重复图id {dup_drawing_ids}')
            if double_titles:
                failures.append(f'{path.parent.name}: 双前缀表题 {double_titles}')

            sections = chapter_sections(text)
            section_map = {num: sec for num, _title, sec in sections}
            required = expectations[paper_type]
            for chapter_num, need in required.items():
                sec = section_map.get(chapter_num, '')
                if not sec:
                    failures.append(f'{path.parent.name}: 缺少第{chapter_num}章正文')
                    continue
                counts = {
                    'tables': len(re.findall(r'<table\\b[^>]*>', sec)),
                    'charts': len(re.findall(r'<chart\\b[^>]*>', sec)),
                    'drawings': len(re.findall(r'<drawing\\b[^>]*>', sec)),
                }
                for kind, minimum in need.items():
                    if counts.get(kind, 0) < minimum:
                        failures.append(
                            f'{path.name}: 第{chapter_num}章 {kind}不足 '
                            f'({counts.get(kind, 0)} < {minimum})'
                        )

        if failures:
            raise SystemExit('\\n'.join(failures))
        """
    ).strip()
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def check_title_propagation(stub_root: Path) -> tuple[bool, str]:
    code = textwrap.dedent(
        """
        import tempfile
        from pathlib import Path
        import generator
        from core import txt_to_docx_safe, extract_docx_text

        def fake_llm(prompt, max_tokens=0):
            if '标签格式校验专家' in prompt:
                return '{"issues_found": false, "corrections": []}'
            if '参考文献' in prompt and '列出' in prompt:
                return '[1] 张三. 示例论文[D]. 某校, 2024.'
            if '摘要' in prompt and '关键词' in prompt:
                return '本文围绕示例对象展开研究。\\n关键词：示例；测试'
            if '第1章' in prompt or '第2章' in prompt or '第3章' in prompt or '第4章' in prompt or '第5章' in prompt or '第6章' in prompt:
                title = prompt.split('第', 1)[-1].split('章', 1)[-1].split('\\n', 1)[0].strip()
                return f'第1章 {title}\\n1.1 内容\\n结论性文字。'
            return '默认内容'

        old_llm = generator.call_llm
        generator.call_llm = fake_llm
        try:
            samples = [
                ('管理', {'title': 'A公司财务管理问题与对策研究', 'company': 'A公司', 'major': '工商管理'}),
                ('设计', {'title': '新中式女装改良设计', 'design_type': '服装设计', 'design_object': '新中式女装'}),
                ('机械', {'title': '某夹具设计与分析', 'mech_object': '铣削夹具', 'major': '机械设计'}),
                ('土木', {'title': '某市办公楼建筑与结构设计', 'object_name': '某市办公楼', 'major': '土木工程'}),
            ]
            for paper_type, profile in samples:
                txt = generator.generate(profile, paper_type)
                if not txt.startswith(f'论文题目：{profile[\"title\"]}'):
                    raise SystemExit(f'{paper_type} txt missing title header')

            with tempfile.TemporaryDirectory() as td:
                txt_path = Path(td) / 'paper.txt'
                docx_path = Path(td) / 'paper.docx'
                txt_path.write_text('论文题目：示例论文题目\\n\\n摘要\\n示例摘要\\n\\n关键词\\n示例\\n', encoding='utf-8')
                txt_to_docx_safe(str(txt_path), str(docx_path), update=lambda *_: None, cover_info=None, drawing_images={})
                docx_text = extract_docx_text(str(docx_path))
                if '论文题目：示例论文题目' not in docx_text:
                    raise SystemExit('docx cover missing title')
        finally:
            generator.call_llm = old_llm
        """
    ).strip()
    proc = run_python_real([], code=code)
    if proc.returncode != 0:
        return False, f"exit={proc.returncode}; stderr={proc.stderr.strip()}"
    return True, "ok"


def main() -> int:
    results: list[tuple[str, bool, str]] = []
    with tempfile.TemporaryDirectory() as tmp:
        stub_root = Path(tmp) / "stubs"
        stub_root.mkdir(parents=True, exist_ok=True)
        write_stub_modules(stub_root)

        record(results, "delivery.batch --help", *check_help(stub_root, "delivery.batch", ["--workflow", "--input", "--output-dir"]))
        record(results, "delivery.audit --help", *check_help(stub_root, "delivery.audit"))
        record(results, "profile --help", *check_profile_help(stub_root))
        record(results, "batch_generate --help", *check_batch_help(stub_root))
        record(results, "renderer --help", *check_renderer_help(stub_root))
        record(results, "delivery.service import", *check_import_service(stub_root))
        record(results, "normalize_paper_type mapping", *check_normalize_mapping(stub_root))
        record(results, "batch arg parsing", *check_batch_parsing(stub_root))
        record(results, "file/path conventions", *check_path_conventions(stub_root))
        record(results, "table matrix guard", *check_table_matrix_guard(stub_root))
        record(results, "multiline table parse", *check_multiline_table_parse(stub_root))
        record(results, "cover aliases", *check_cover_aliases(stub_root))
        record(results, "multiline chart parse", *check_multiline_chart_parse(stub_root))
        record(results, "multiline drawing parse", *check_multiline_drawing_parse(stub_root))
        record(results, "partial drawing parse", *check_partial_drawing_parse(stub_root))
        record(results, "placeholder drawing title replace", *check_placeholder_drawing_title_replace(stub_root))
        record(results, "mech json merge", *check_mech_json_merge(stub_root))
        record(results, "manage markup fallback", *check_manage_markup_fallback(stub_root))
        record(results, "manage markup fallback with later table", *check_manage_markup_fallback_with_later_table(stub_root))
        record(results, "annotation wrapper flatten", *check_annotation_wrapper_flatten(stub_root))
        record(results, "civil generation", *check_civil_generation(stub_root))
        record(results, "civil markup normalization", *check_civil_markup_normalization(stub_root))
        record(results, "civil label refinement", *check_civil_label_refinement(stub_root))
        record(results, "civil batch/render support", *check_civil_batch_and_render_support(stub_root))
        record(results, "real sample structure guard", *check_real_sample_structure_guard(stub_root))
        record(results, "title propagation", *check_title_propagation(stub_root))
        record(results, "table tag order valid", *check_table_tag_order_valid(stub_root))

    passed = 0
    failed = 0
    for name, ok, detail in results:
        if ok:
            passed += 1
            print(f"PASS {name}")
        else:
            failed += 1
            print(f"FAIL {name}: {detail}")
    print(f"SUMMARY {passed} passed, {failed} failed")
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
