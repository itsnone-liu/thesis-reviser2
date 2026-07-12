#!/usr/bin/env python3
"""优化部分标题中"以"后面的内容，只保留公司/对象名称。"""
import os, re
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import openpyxl

BASE = "/root/project/workspace/论文资料/整理"
XLSX = f"{BASE}/学位申请合并_含论文状态.xlsx"

# 需要优化标题的同学
OPTIMIZE = {
    # 准考证: (新标题, 原因)
    "29820430393": ("旅游管理问题研究-以无锡太湖学院为例", "正确"),
    "29822440011": ("融资风险问题研究-以格力电器为例", "去掉过长修饰"),
    "29822440007": ("营运资金管理问题研究-以南方黑芝麻集团为例", "去掉修饰"),
    "29822440002": ("财务风险控制问题研究-以海底捞为例", "简化为公司名"),
    "29823430793": ("培训体系问题研究-以M数据公司为例", "简化为公司名"),
    "29822430644": ("信贷风险管理问题研究-以中信银行为例", "修正格式"),
    "29823430177": ("培训体系优化问题研究-以世茂假日酒店为例", "简化"),
    "29822440016": ("现金流管理问题研究-以蔚来汽车为例", "统一格式"),
}

wb = openpyxl.load_workbook(XLSX)
ws = wb['学位申请汇总']

for r in range(2, ws.max_row + 1):
    sid_raw = ws.cell(r, 2).value
    if sid_raw is None:
        continue
    sid = str(int(sid_raw))
    
    if sid in OPTIMIZE:
        name = ws.cell(r, 1).value
        new_title, reason = OPTIMIZE[sid]
        
        # 找到论文文件修改封面
        found = False
        for batch in ["25", "26"]:
            for major in ["经管", "机械", "设计"]:
                d = os.path.join(BASE, "论文终版", batch, major)
                if not os.path.exists(d):
                    continue
                for f in os.listdir(d):
                    if sid in f and f.endswith('.docx'):
                        filepath = os.path.join(d, f)
                        doc = Document(filepath)
                        for p in doc.paragraphs:
                            if '论文题目' in p.text:
                                for run in p.runs:
                                    if '论文题目' in run.text:
                                        run.text = f"论文题目：{new_title}"
                                    else:
                                        run.text = ''
                                break
                        doc.save(filepath)
                        found = True
                        break
                if found:
                    break
            if found:
                break
        
        # 更新状态表
        ws.cell(r, 6).value = new_title
        
        print(f"  ✅ {name}: {new_title} ({reason})")

wb.save(XLSX)
print(f"\n✅ 全部优化完成并已保存")
