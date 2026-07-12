#!/usr/bin/env python3
"""
环艺设计类论文批量修改脚本 — 直接修复DOCX
针对审核表中的低/无等级论文进行批量修复

主要修复：
1. 删摘要"摘要："
2. 重建表格（从TXT解析HTML风格标签）
3. 封面信息修正（层次、学号、专业）
"""

import os
import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import openpyxl

# ==================== 配置 ====================
HUANYI_DIR = Path("/root/project/workspace/论文资料/整理/环艺")
MODIFIED_DIR = Path("/root/project/workspace/论文资料/整理/设计类_修改版")
OUTPUT_DIR = Path("/root/project/workspace/论文资料/整理/环艺_修改版")
AUDIT_FILE = Path("/root/project/workspace/论文资料/整理/设计类审核表.xlsx")

# ==================== 工具函数 ====================

def set_run_font(run, font_name="宋体", size=12, bold=False):
    run.font.name = font_name
    try:
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    except:
        pass
    run.font.size = Pt(size)
    run.font.bold = bold


def set_cell_font(cell, font_name="宋体", size=9, bold=False, align_center=True):
    """设置单元格字体和对齐"""
    for para in cell.paragraphs:
        if align_center:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in para.runs:
            set_run_font(run, font_name, size, bold)


def find_docx_file(exam_id):
    """在环艺目录中找到对应准考证号的DOCX文件"""
    for f in os.listdir(HUANYI_DIR):
        if exam_id in f and f.endswith('.docx'):
            return HUANYI_DIR / f
    return None


def find_txt_file(exam_id):
    """在设计类_修改版中找到对应的TXT文件"""
    # 精准匹配
    for d in os.listdir(MODIFIED_DIR):
        if exam_id in d and d.startswith(exam_id[:4]):
            txt_path = MODIFIED_DIR / d / "paper.txt"
            if txt_path.exists():
                return txt_path
    # 模糊匹配（从后6位）
    short = exam_id[-6:]
    for d in os.listdir(MODIFIED_DIR):
        if short in d:
            txt_path = MODIFIED_DIR / d / "paper.txt"
            if txt_path.exists():
                return txt_path
    return None


# ==================== TXT表格解析 ====================

def extract_tables_from_txt(txt_path):
    """
    从TXT提取表格数据。设计类表格格式：
    <table />
     <caption标题</caption
     <tr
     <th列1</th
     <th列2</th
     </tr
     <tr
     <td数据1</td
     <td数据2</td
     </tr
    </table
    返回 [(标题, 表头列表, 数据行列表), ...]
    """
    with open(txt_path, 'r') as f:
        lines = f.readlines()
    
    tables = []
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        # 找到表格开始
        if stripped.startswith('<table') or stripped == '<table':
            # 收集此表格的所有行
            table_lines = [lines[i]]
            i += 1
            while i < len(lines):
                s = lines[i].strip()
                if s.startswith('</table') or s == '</table' or re.match(r'</table\s*$', s):
                    table_lines.append(lines[i])
                    i += 1
                    break
                table_lines.append(lines[i])
                i += 1
            # 解析表格
            table_data = parse_html_table(table_lines)
            if table_data:
                tables.append(table_data)
        else:
            i += 1
    
    return tables


def parse_html_table(table_lines):
    """解析HTML风格表格行，返回 (标题, 表头, 数据行)"""
    caption = ""
    all_rows = []  # [(is_header, [cells])]
    current_row_cells = []
    current_is_header = False
    in_row = False
    row_depth = 0
    
    # 合并所有行为一个字符串
    flat = ''.join(table_lines)
    
    # 提取caption
    cap_m = re.search(r'<caption[^>]*>([^<]*)</caption', flat)
    if cap_m:
        caption = cap_m.group(1).strip()
    
    # 按行解析
    for line in table_lines:
        s = line.strip()
        
        if '<tr' in s or s.startswith('<tr'):
            # 新行开始
            if current_row_cells:
                all_rows.append((current_is_header, current_row_cells))
            current_row_cells = []
            current_is_header = ('<th' in s)
            in_row = True
            # 提取当前行中的内容
            cells = re.findall(r'<(?:th|td)\b[^>]*>(.*?)</(?:th|td)\s*>', s, re.DOTALL)
            if cells:
                current_row_cells.extend([c.strip() for c in cells])
        
        elif '</tr' in s or s == '</tr':
            if current_row_cells:
                all_rows.append((current_is_header, current_row_cells))
            current_row_cells = []
            current_is_header = False
            in_row = False
        
        elif '<th' in s or '<td' in s:
            # 可能是独立的行，没有<tr包装（如张颖的格式）
            cells = re.findall(r'<(?:th|td)\b[^>]*>(.*?)</(?:th|td)\s*>', s, re.DOTALL)
            is_h = '<th' in s
            if cells:
                if is_h:
                    all_rows.append((True, [c.strip() for c in cells]))
                else:
                    all_rows.append((False, [c.strip() for c in cells]))
    
    # 最后一行
    if current_row_cells:
        all_rows.append((current_is_header, current_row_cells))
    
    if not all_rows:
        return None
    
    # 判断表头：如果有th行，第一行为表头
    headers = []
    rows_data = []
    
    for is_header, cells in all_rows:
        if is_header and not headers:
            headers = cells
        elif is_header and headers:
            # 第二个th行也当表头
            headers.extend(cells)
        elif not is_header:
            if not headers:
                headers = cells
            else:
                # 确保列数匹配
                if len(cells) >= len(headers):
                    rows_data.append(cells[:len(headers)])
                else:
                    # 补全
                    padded = cells + [''] * (len(headers) - len(cells))
                    rows_data.append(padded)
    
    # 处理一种特殊格式：如果某行同时有th和td
    # 以及只有数据没有表头的情况
    if not headers and all_rows:
        for _, cells in all_rows:
            rows_data.append(cells)
    
    return (caption, headers, rows_data)


def validate_table_data(tables):
    """打印表格数据验证"""
    for ti, (caption, headers, rows_data) in enumerate(tables):
        print(f"  表{ti+1}: caption='{caption[:30] if caption else ''}', "
              f"headers={len(headers)}列, data={len(rows_data)}行")
        if headers:
            print(f"    表头: {headers}")
        for ri, row in enumerate(rows_data[:3]):
            print(f"    行{ri+1}: {row}")


# ==================== DOCX表格替换 ====================

def replace_bad_tables_in_docx(doc, docx_path, tables_from_txt):
    """
    替换DOCX中的格式错误表格（1列多行的表格）。
    策略：找到DOCX中所有1列的表格，按顺序用TXT表格替换。
    """
    # 找到所有1列表格
    bad_tables = []
    for ti, table in enumerate(doc.tables):
        if len(table.columns) == 1 and len(table.rows) >= 3:
            bad_tables.append(ti)
    
    if not bad_tables:
        return "无1列表格可替换"
    
    # 记录表格起始的p元素位置（用于定位替换）
    body = doc.part.document.body._element
    all_p_elements = list(body.findall(qn('w:p')))
    
    replaced_count = 0
    txt_table_idx = 0
    
    for ti in bad_tables:
        if txt_table_idx >= len(tables_from_txt):
            break
        
        table = doc.tables[ti]
        caption, headers, rows_data = tables_from_txt[txt_table_idx]
        
        if not headers or len(headers) < 2:
            txt_table_idx += 1
            continue
        
        # 替换表格
        _rebuild_table_elements(doc, table, caption, headers, rows_data)
        replaced_count += 1
        txt_table_idx += 1
    
    return f"替换{replaced_count}个表格"


def _rebuild_table_elements(doc, table, caption, headers, rows_data):
    """
    重建一个python-docx表格对象的内容。
    保持表格对象不变，只修改其内部的XML结构。
    """
    num_cols = len(headers)
    num_rows = 1 + len(rows_data)
    
    # 获取表格的tbl元素
    tbl = table._tbl
    ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    
    # 找到所有行，删除多余的行或增加行
    existing_rows = list(tbl.findall(f'{{{ns}}}tr'))
    
    while len(existing_rows) > num_rows:
        tbl.remove(existing_rows[-1])
        existing_rows = list(tbl.findall(f'{{{ns}}}tr'))
    
    while len(existing_rows) < num_rows:
        # 复制最后一行
        from lxml import etree
        new_row = etree.SubElement(tbl, f'{{{ns}}}tr')
        last_row = existing_rows[-1]
        for tc in last_row.findall(f'{{{ns}}}tc'):
            new_tc = etree.SubElement(new_row, f'{{{ns}}}tc')
            new_p = etree.SubElement(new_tc, f'{{{ns}}}p')
            new_r = etree.SubElement(new_p, f'{{{ns}}}r')
            new_t = etree.SubElement(new_r, f'{{{ns}}}t')
            new_t.text = ''
        existing_rows = list(tbl.findall(f'{{{ns}}}tr'))
    
    # 现在填充数据
    from lxml import etree as _etree
    
    # 表头行
    tr = existing_rows[0]
    tcs = list(tr.findall(f'{{{ns}}}tc'))
    
    # 调整表头列数
    while len(tcs) > num_cols:
        tr.remove(tcs[-1])
        tcs = list(tr.findall(f'{{{ns}}}tc'))
    while len(tcs) < num_cols:
        new_tc = _etree.SubElement(tr, f'{{{ns}}}tc')
        new_p = _etree.SubElement(new_tc, f'{{{ns}}}p')
        new_r = _etree.SubElement(new_p, f'{{{ns}}}r')
        new_t = _etree.SubElement(new_r, f'{{{ns}}}t')
        new_t.text = ''
        tcs = list(tr.findall(f'{{{ns}}}tc'))
    
    # 填充表头
    for j, h in enumerate(headers):
        if j < len(tcs):
            tc = tcs[j]
            # 清空
            for p in list(tc.findall(f'{{{ns}}}p')):
                tc.remove(p)
            # 新建段落
            new_p = _etree.SubElement(tc, f'{{{ns}}}p')
            pPr = _etree.SubElement(new_p, f'{{{ns}}}pPr')
            jc = _etree.SubElement(pPr, f'{{{ns}}}jc')
            jc.set(f'{{{ns}}}val', 'center')
            new_r = _etree.SubElement(new_p, f'{{{ns}}}r')
            rPr = _etree.SubElement(new_r, f'{{{ns}}}rPr')
            rFonts = _etree.SubElement(rPr, f'{{{ns}}}rFonts')
            rFonts.set(f'{{{ns}}}ascii', '宋体')
            rFonts.set(f'{{{ns}}}eastAsia', '宋体')
            sz = _etree.SubElement(rPr, f'{{{ns}}}sz')
            sz.set(f'{{{ns}}}val', '18')  # 9pt = 18 half-points
            b = _etree.SubElement(rPr, f'{{{ns}}}b')
            new_t = _etree.SubElement(new_r, f'{{{ns}}}t')
            new_t.text = h
            # 设置列宽
            tcPr = tc.find(f'{{{ns}}}tcPr')
            if tcPr is None:
                tcPr = _etree.SubElement(tc, f'{{{ns}}}tcPr')
            tcW = tcPr.find(f'{{{ns}}}tcW')
            if tcW is None:
                tcW = _etree.SubElement(tcPr, f'{{{ns}}}tcW')
            tcW.set(f'{{{ns}}}w', str(int(5000 / num_cols)))
            tcW.set(f'{{{ns}}}type', 'pct')
    
    # 数据行
    for i, row_data in enumerate(rows_data):
        if i + 1 >= len(existing_rows):
            break
        tr = existing_rows[i + 1]
        tcs = list(tr.findall(f'{{{ns}}}tc'))
        
        # 调整列数
        while len(tcs) > num_cols:
            tr.remove(tcs[-1])
            tcs = list(tr.findall(f'{{{ns}}}tc'))
        while len(tcs) < num_cols:
            new_tc = _etree.SubElement(tr, f'{{{ns}}}tc')
            new_p = _etree.SubElement(new_tc, f'{{{ns}}}p')
            new_r = _etree.SubElement(new_p, f'{{{ns}}}r')
            new_t = _etree.SubElement(new_r, f'{{{ns}}}t')
            new_t.text = ''
            tcs = list(tr.findall(f'{{{ns}}}tc'))
        
        for j, val in enumerate(row_data):
            if j < len(tcs):
                tc = tcs[j]
                for p in list(tc.findall(f'{{{ns}}}p')):
                    tc.remove(p)
                new_p = _etree.SubElement(tc, f'{{{ns}}}p')
                pPr = _etree.SubElement(new_p, f'{{{ns}}}pPr')
                jc = _etree.SubElement(pPr, f'{{{ns}}}jc')
                jc.set(f'{{{ns}}}val', 'center')
                new_r = _etree.SubElement(new_p, f'{{{ns}}}r')
                rPr = _etree.SubElement(new_r, f'{{{ns}}}rPr')
                rFonts = _etree.SubElement(rPr, f'{{{ns}}}rFonts')
                rFonts.set(f'{{{ns}}}ascii', '宋体')
                rFonts.set(f'{{{ns}}}eastAsia', '宋体')
                sz = _etree.SubElement(rPr, f'{{{ns}}}sz')
                sz.set(f'{{{ns}}}val', '18')
                new_t = _etree.SubElement(new_r, f'{{{ns}}}t')
                new_t.text = val
    
    # 设置表格宽100%
    tblPr = tbl.find(f'{{{ns}}}tblPr')
    if tblPr is None:
        from lxml import etree as _etree2
        tblPr = _etree2.Element(f'{{{ns}}}tblPr')
        tbl.insert(0, tblPr)
    tblW = tblPr.find(f'{{{ns}}}tblW')
    if tblW is None:
        from lxml import etree as _etree3
        tblW = _etree3.SubElement(tblPr, f'{{{ns}}}tblW')
    tblW.set(f'{{{ns}}}w', '5000')
    tblW.set(f'{{{ns}}}type', 'pct')


# ==================== 摘要修复 ====================

def fix_abstract_prefix(doc):
    """
    删除摘要正文段落开头的"摘要："或"摘要:"
    只操作run内的文字，完全保留格式
    """
    fixes = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        # 找"摘要："开头的正文段落（不是单独的摘要标题）
        if not (text.startswith('摘要：') or text.startswith('摘要:')):
            continue
        
        # 确认这是正文段不是标题
        if text == '摘要：':
            continue
        
        prefix = '摘要：' if text.startswith('摘要：') else '摘要:'
        
        # 在run级别删除前缀
        remaining = len(prefix)
        for run in p.runs:
            if not run.text:
                continue
            if remaining <= 0:
                break
            run_len = len(run.text)
            if run_len <= remaining:
                remaining -= run_len
                run.text = ''
            else:
                run.text = run.text[remaining:]
                remaining = 0
        
        fixes += 1
    
    return fixes


# ==================== 封面修复 ====================

def fix_cover(doc, exam_id, problems):
    """修复封面信息"""
    fixes = []
    
    # 问题分类
    has_level_fix = '层次' in problems and '专升本' in problems
    has_major_fix = '专业名称不对' in problems or '环境设计' in problems
    has_id_fix = '学号' in problems
    
    # 提取正确的学号
    correct_id = exam_id
    m = re.search(r'应为[：:"\u201c]*(\d+)', problems)
    if m:
        correct_id = m.group(1)
    
    for p in doc.paragraphs:
        text = ''.join(r.text for r in p.runs)
        
        # 学号
        if has_id_fix and ('学    号' in text or '学号' in text):
            for run in p.runs:
                if '学' in run.text or '号' in run.text:
                    # 替换学号
                    m2 = re.search(r'(\d{6,})', run.text)
                    if m2:
                        orig = m2.group(1)
                        if orig != correct_id:
                            run.text = run.text.replace(orig, correct_id)
                            fixes.append(f"学号: {orig}→{correct_id}")
                    break
        
        # 层次（本科→专升本）
        if has_level_fix:
            if '层' in text and '次' in text:
                for run in p.runs:
                    if '本科' in run.text:
                        run.text = run.text.replace('本科', '专升本')
                        fixes.append("层次: 本科→专升本")
                    break
        
        # 专业
        if has_major_fix and ('专' in text and '业' in text):
            if '环境设计' not in text:
                for run in p.runs:
                    if '环境' not in run.text:
                        run.text = re.sub(
                            r'(专\s*业\s*[：:])\s*\S+',
                            r'\1环境设计',
                            run.text
                        )
                        fixes.append("专业→环境设计")
                        break
    
    return fixes


# ==================== 其他修复 ====================

def fix_other_issues(doc, problems):
    """处理其他杂项问题"""
    fixes = []
    
    # 修复"本研究团队"→"本研究"
    if '本研究团队' in problems:
        for p in doc.paragraphs:
            for run in p.runs:
                if '本研究团队' in run.text:
                    run.text = run.text.replace('本研究团队', '本研究')
                    fixes.append("本研究团队→本研究")
    
    # 修复缺少学号
    if '缺少学号' in problems:
        m = re.search(r'应为[：:"\u201c]*(\d+)', problems)
        if m:
            correct_id = m.group(1)
            for p in doc.paragraphs:
                text = ''.join(r.text for r in p.runs)
                if '学    号' in text or '学号' in text:
                    for run in p.runs:
                        if run.text.strip() == '' or (run.text.strip() == '学    号：'):
                            continue
                        m2 = re.search(r'(\d{6,})', run.text)
                        if not m2 or m2.group(1) != correct_id:
                            # 替换或添加学号
                            pass
    
    return fixes


# ==================== 目录页码修复 ====================

def fix_toc_pages(doc):
    """修正目录中的页码（重新计算）"""
    # 目录页码修复：通常需要在渲染时处理
    # 在DOCX层面改动目录页码需要操作XML域代码
    # 这是渲染层面的问题，DOCX修改较复杂
    return []


# ==================== 主处理函数 ====================

def process_single(exam_id, name, problems):
    """处理单篇论文"""
    # 1. 找到DOCX
    docx_path = find_docx_file(exam_id)
    if not docx_path:
        return {'status': 'error', 'reason': f'找不到文件 {exam_id}'}
    
    print(f"  读取: {docx_path.name}")
    
    # 2. 找到TXT（用于表格数据）
    txt_path = find_txt_file(exam_id)
    
    # 3. 读取DOCX
    doc = Document(str(docx_path))
    
    report = {'fixes': []}
    
    # 4. 逐一修复
    # 4.1 删摘要前缀
    if '摘要部分' in problems and ('多了' in problems or '摘要' in problems):
        n = fix_abstract_prefix(doc)
        if n:
            report['fixes'].append(f"删摘要前缀: {n}处")
    
    # 4.2 修复表格
    if ('表格格式不对' in problems or '表格内容不对' in problems) and txt_path:
        tables = extract_tables_from_txt(txt_path)
        if tables:
            print(f"  TXT中提取到 {len(tables)} 个表格")
            result = replace_bad_tables_in_docx(doc, docx_path, tables)
            report['fixes'].append(f"表格: {result}")
        else:
            report['fixes'].append("表格: TXT中无表格数据")
    
    # 4.3 封面修复
    cover_fixes = fix_cover(doc, exam_id, problems)
    if cover_fixes:
        report['fixes'].extend(cover_fixes)
    
    # 4.4 其他修复
    other_fixes = fix_other_issues(doc, problems)
    if other_fixes:
        report['fixes'].extend(other_fixes)
    
    # 5. 保存到输出目录
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_name = f"{exam_id}_{name}_修改版.docx"
    output_path = OUTPUT_DIR / output_name
    doc.save(str(output_path))
    
    report['output'] = str(output_path)
    return report


def main():
    # 读取审核表
    wb = openpyxl.load_workbook(str(AUDIT_FILE))
    ws = wb['Sheet1']
    
    papers = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name, exam_id_raw, _, problems, level = row
        exam_id = str(exam_id_raw or '').strip().zfill(12)
        level_str = str(level).strip() if level else 'None'
        if level_str not in ('低', 'None', 'None'):
            continue
        papers.append({
            'name': name or '未知',
            'exam_id': exam_id,
            'problems': problems or '',
        })
    
    # 命令行筛选
    if len(sys.argv) > 1:
        target_ids = sys.argv[1:]
        papers = [p for p in papers if any(tid in p['exam_id'] for tid in target_ids)]
    
    print(f"需要处理的论文: {len(papers)}篇\n")
    
    results = []
    for i, paper in enumerate(papers):
        name = paper['name']
        exam_id = paper['exam_id']
        problems = paper['problems']
        
        print(f"[{i+1}/{len(papers)}] {name} ({exam_id})")
        if problems:
            for line in problems.split('\n'):
                if line.strip():
                    print(f"  📋 {line.strip()[:80]}")
        
        result = process_single(exam_id, name, problems)
        results.append(result)
        
        if result['status'] == 'error':
            print(f"  ❌ {result['reason']}")
        else:
            for fix in result['fixes']:
                print(f"  ✅ {fix}")
            if result.get('output'):
                print(f"  📄 {result['output']}")
        print()
    
    success = sum(1 for r in results if r['status'] != 'error')
    failed = sum(1 for r in results if r['status'] == 'error')
    print(f"\n完成！成功 {success}/{len(papers)}, 失败 {failed}")


if __name__ == '__main__':
    main()
