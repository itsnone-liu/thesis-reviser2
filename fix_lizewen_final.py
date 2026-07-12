#!/usr/bin/env python3
"""
李泽文最终修复 — 从原始文件直接修改
只做必要修改，不引入之前v2的错误
"""

import os, re, zipfile, shutil
from pathlib import Path
from lxml import etree
from docx import Document

NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
HUANYI_DIR = Path("/root/project/workspace/论文资料/整理/环艺")
OUTPUT_DIR = Path("/root/project/workspace/论文资料/整理/环艺_修改版")

def make_table(headers, data_rows):
    num_cols = len(headers)
    tbl = etree.Element(f'{{{NS}}}tbl')
    tblPr = etree.SubElement(tbl, f'{{{NS}}}tblPr')
    etree.SubElement(tblPr, f'{{{NS}}}tblStyle').set(f'{{{NS}}}val', 'Table Grid')
    tw = etree.SubElement(tblPr, f'{{{NS}}}tblW')
    tw.set(f'{{{NS}}}w', '5000')
    tw.set(f'{{{NS}}}type', 'pct')
    tblGrid = etree.SubElement(tbl, f'{{{NS}}}tblGrid')
    for ci in range(num_cols):
        gc = etree.SubElement(tblGrid, f'{{{NS}}}gridCol')
        gc.set(f'{{{NS}}}w', str(int(5000/num_cols)))
    
    def cell(text, bold=False):
        tc = etree.Element(f'{{{NS}}}tc')
        p = etree.SubElement(tc, f'{{{NS}}}p')
        pp = etree.SubElement(p, f'{{{NS}}}pPr')
        etree.SubElement(pp, f'{{{NS}}}jc').set(f'{{{NS}}}val', 'center')
        r = etree.SubElement(p, f'{{{NS}}}r')
        rp = etree.SubElement(r, f'{{{NS}}}rPr')
        rf = etree.SubElement(rp, f'{{{NS}}}rFonts')
        rf.set(f'{{{NS}}}ascii', '宋体')
        rf.set(f'{{{NS}}}eastAsia', '宋体')
        etree.SubElement(rp, f'{{{NS}}}sz').set(f'{{{NS}}}val', '18')
        if bold:
            etree.SubElement(rp, f'{{{NS}}}b')
        t = etree.SubElement(r, f'{{{NS}}}t')
        t.text = text if text else ''
        return tc
    
    tr = etree.SubElement(tbl, f'{{{NS}}}tr')
    for h in headers:
        tr.append(cell(h, bold=True))
    for row in data_rows:
        tr = etree.SubElement(tbl, f'{{{NS}}}tr')
        for val in row:
            tr.append(cell(val))
    return tbl


# 1. 复制原始
src = HUANYI_DIR / "029821410170_修改版.docx"
out = OUTPUT_DIR / "029821410170_李泽文_修改版.docx"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
shutil.copy2(str(src), str(out))
print("✅ 已复制原始文件")

# 2. python-docx修改摘要和学号
doc = Document(str(out))

# 删摘要
n = 0
for p in doc.paragraphs:
    text = p.text.strip()
    if text.startswith('摘要：') and text != '摘要：':
        remaining = 3
        for run in p.runs:
            if not run.text or remaining <= 0:
                continue
            if len(run.text) <= remaining:
                remaining -= len(run.text)
                run.text = ''
            else:
                run.text = run.text[remaining:]
                remaining = 0
        n += 1
print(f"✅ 删摘要前缀 {n}处")

# 学号
id_ok = False
for p in doc.paragraphs:
    full = ''.join(r.text for r in p.runs)
    if '学    号' in full:
        for run in p.runs:
            m = re.search(r'(\d{6,})', run.text)
            if m and m.group(1) != '029821410170':
                run.text = run.text.replace(m.group(1), '029821410170')
                id_ok = True
                break
if id_ok:
    print("✅ 学号→029821410170")

doc.save(str(out))

# 3. lxml修改表格结构
with zipfile.ZipFile(str(out), 'r') as z:
    doc_xml = etree.fromstring(z.read('word/document.xml'))
    others = {n: z.read(n) for n in z.namelist() if n != 'word/document.xml'}

body = doc_xml.find(f'{{{NS}}}body')

# 3a. 修复3.2节表格的tblGrid
for child in list(body):
    if child.tag == f'{{{NS}}}tbl':
        tblGrid = child.find(f'{{{NS}}}tblGrid')
        if tblGrid is not None:
            gcs = list(tblGrid.findall(f'{{{NS}}}gridCol'))
            if len(gcs) == 1:
                for gc in gcs:
                    tblGrid.remove(gc)
                for _ in range(4):
                    etree.SubElement(tblGrid, f'{{{NS}}}gridCol').set(f'{{{NS}}}w', '1250')
                # 修复每个tr
                for tr in child.findall(f'{{{NS}}}tr'):
                    tcs = list(tr.findall(f'{{{NS}}}tc'))
                    while len(tcs) > 4:
                        tr.remove(tcs[-1])
                        tcs = list(tr.findall(f'{{{NS}}}tc'))
                print("✅ 3.2节表格tblGrid修复")
        break

# 3b. 找Markdown表格并替换
markdown_paras = []
for child in list(body):
    if child.tag == f'{{{NS}}}p':
        texts = ''.join(t.text or '' for t in child.iter(f'{{{NS}}}t') if t.tag == f'{{{NS}}}t')
        text = texts.strip()
        if text.startswith('|') and any(kw in text for kw in ['问题类型', '---', '主题文化', '环保与舒适', '空间功能']):
            markdown_paras.append(child)

if markdown_paras:
    first = markdown_paras[0]
    idx = list(body).index(first)
    
    # 表标题
    title_p = etree.Element(f'{{{NS}}}p')
    pp = etree.SubElement(title_p, f'{{{NS}}}pPr')
    etree.SubElement(pp, f'{{{NS}}}jc').set(f'{{{NS}}}val', 'center')
    r = etree.SubElement(title_p, f'{{{NS}}}r')
    rp = etree.SubElement(r, f'{{{NS}}}rPr')
    rf = etree.SubElement(rp, f'{{{NS}}}rFonts')
    rf.set(f'{{{NS}}}ascii', '宋体')
    rf.set(f'{{{NS}}}eastAsia', '宋体')
    etree.SubElement(rp, f'{{{NS}}}sz').set(f'{{{NS}}}val', '20')
    etree.SubElement(rp, f'{{{NS}}}b')
    etree.SubElement(r, f'{{{NS}}}t').text = '表4.1 问题类型、实例与改进方向对应关系'
    body.insert(idx, title_p)
    
    # 表格
    new_tbl = make_table(
        ['问题类型', '典型实例', '改进方向（非解决方案，仅梳理思路）'],
        [
            ['主题文化表达流于表面', '某"农耕主题"民宿：仅放置旧农具，空间缺乏整体叙事', '强化空间叙事连贯性，材质与主题统一'],
            ['环保与舒适度冲突', '西南夯土民宿：未做内保温，湿度超75%，被褥发霉', '平衡传统外观与现代保温技术'],
            ['空间功能与需求错位', '观景书吧：书籍少、座椅硬，使用率不足30%', '基于游客行为模式重新规划功能'],
        ]
    )
    body.insert(idx, new_tbl)
    
    for elem in markdown_paras:
        body.remove(elem)
    
    print(f"✅ Markdown表格→docx表格（表4.1，删{len(markdown_paras)}行）")
else:
    print("⚠️ 找不到Markdown表格")

# 写回
with zipfile.ZipFile(str(out), 'w', zipfile.ZIP_DEFLATED) as z:
    z.writestr('word/document.xml',
        etree.tostring(doc_xml, xml_declaration=True, encoding='UTF-8', standalone=True))
    for name, data in others.items():
        z.writestr(name, data)

print(f"📄 {out}")
