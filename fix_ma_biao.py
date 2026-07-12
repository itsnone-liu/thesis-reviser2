#!/usr/bin/env python3
"""去掉指导教师马彪名字后的括号内容。"""
import os, re
from docx import Document
import openpyxl

BASE = "/root/project/workspace/论文资料/整理"
XLSX = f"{BASE}/学位申请合并_含论文状态.xlsx"

# 需要处理的 (姓名, 准考证)
TARGETS = [
    ("余心伟", "29822430998"),
    ("吴策", "29822430988"),
    ("侍继典", "29822430978"),
]

def find_file(sid):
    for batch in ["25", "26"]:
        for major in ["经管", "机械", "设计"]:
            d = os.path.join(BASE, "论文终版", batch, major)
            if not os.path.exists(d):
                continue
            for f in os.listdir(d):
                if sid in f and f.endswith('.docx'):
                    return os.path.join(d, f)
    return None

# 1. 改封面
for name, sid in TARGETS:
    fp = find_file(sid)
    if not fp:
        print(f"  ⚠ {name}: 未找到文件")
        continue
    
    doc = Document(fp)
    modified = False
    
    for p in doc.paragraphs:
        text = p.text.strip()
        if '指导教师' in text or '指导老师' in text:
            for r in p.runs:
                if '马彪' in r.text:
                    r.text = '指导教师：马彪'
                    modified = True
                elif '职称' in r.text or '(' in r.text or '（' in r.text:
                    r.text = ''
    
    if modified:
        doc.save(fp)
        print(f"  ✅ {name}: 封面已修改")
    else:
        print(f"  ⚠ {name}: 未找到需要修改的内容")

# 2. 更新状态表
print()
wb = openpyxl.load_workbook(XLSX)
ws = wb['学位申请汇总']

for r in range(2, ws.max_row + 1):
    teacher = str(ws.cell(r, 7).value or '')
    if '马彪' in teacher and '职称' in teacher:
        name = ws.cell(r, 1).value
        ws.cell(r, 7).value = '马彪'
        print(f"  ✅ {name}: 状态表已更新")

wb.save(XLSX)
print(f"\n✅ 全部完成")
