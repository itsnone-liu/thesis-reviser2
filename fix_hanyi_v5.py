#!/usr/bin/env python3
"""
环艺设计类论文精确修复 v5 — 李泽文 + 冯嘉诚
策略：先用python-docx修改段落/run，保存；再打开用lxml修改表格等XML结构
"""

import os, re
from pathlib import Path
from lxml import etree
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
HUANYI_DIR = Path("/root/project/workspace/论文资料/整理/环艺")
OUTPUT_DIR = Path("/root/project/workspace/论文资料/整理/环艺_修改版")


def find_docx(exam_id):
    for f in os.listdir(HUANYI_DIR):
        if exam_id in f and f.endswith('.docx'):
            return HUANYI_DIR / f
    return None


def make_table(headers, data_rows):
    num_cols = len(headers)
    tbl = etree.Element(f'{{{NS}}}tbl')
    tblPr = etree.SubElement(tbl, f'{{{NS}}}tblPr')
    tblStyle = etree.SubElement(tblPr, f'{{{NS}}}tblStyle')
    tblStyle.set(f'{{{NS}}}val', 'Table Grid')
    tblW = etree.SubElement(tblPr, f'{{{NS}}}tblW')
    tblW.set(f'{{{NS}}}w', '5000')
    tblW.set(f'{{{NS}}}type', 'pct')
    tblGrid = etree.SubElement(tbl, f'{{{NS}}}tblGrid')
    for ci in range(num_cols):
        gc = etree.SubElement(tblGrid, f'{{{NS}}}gridCol')
        gc.set(f'{{{NS}}}w', str(int(5000/num_cols)))
    
    def make_cell(text, bold=False):
        tc = etree.Element(f'{{{NS}}}tc')
        p = etree.SubElement(tc, f'{{{NS}}}p')
        pPr = etree.SubElement(p, f'{{{NS}}}pPr')
        jc = etree.SubElement(pPr, f'{{{NS}}}jc')
        jc.set(f'{{{NS}}}val', 'center')
        r = etree.SubElement(p, f'{{{NS}}}r')
        rPr = etree.SubElement(r, f'{{{NS}}}rPr')
        rf = etree.SubElement(rPr, f'{{{NS}}}rFonts')
        rf.set(f'{{{NS}}}ascii', '宋体')
        rf.set(f'{{{NS}}}eastAsia', '宋体')
        sz = etree.SubElement(rPr, f'{{{NS}}}sz')
        sz.set(f'{{{NS}}}val', '18')
        if bold:
            etree.SubElement(rPr, f'{{{NS}}}b')
        t = etree.SubElement(r, f'{{{NS}}}t')
        t.text = text if text else ''
        return tc
    
    tr = etree.SubElement(tbl, f'{{{NS}}}tr')
    for h in headers:
        tr.append(make_cell(h, bold=True))
    for row in data_rows:
        tr = etree.SubElement(tbl, f'{{{NS}}}tr')
        for val in row:
            tr.append(make_cell(val))
    return tbl


# ==================== 第一步：python-docx修改段落 ====================

def step1_lizewen():
    print("Step 1: 段落修改（摘要前缀 + 封面学号）")
    src = find_docx('029821410170')
    doc = Document(str(src))
    
    # 删摘要
    n = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        if text.startswith('摘要：') and text != '摘要：':
            remaining = 3
            for run in p.runs:
                if not run.text or remaining <= 0:
                    continue
                if len(run.text) <= remaining:
                    remaining -= len(run.text)
                    run.text = ''
                else:
                    run.text = run.text[remaining:]
                    remaining = 0
            n += 1
    print(f"  ✅ 删摘要前缀 {n}处")
    
    # 封面板学号
    id_ok = False
    for p in doc.paragraphs:
        full = ''.join(r.text for r in p.runs)
        if '学    号' in full:
            for run in p.runs:
                m = re.search(r'(\d{6,})', run.text)
                if m and m.group(1) != '029821410170':
                    run.text = run.text.replace(m.group(1), '029821410170')
                    id_ok = True
                    break
    if id_ok:
        print("  ✅ 学号→029821410170")
    
    # 保存临时文件
    tmp = OUTPUT_DIR / "_tmp_lizewen.docx"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(tmp))
    print(f"  临时保存: {tmp}")
    return tmp


def step1_fengjiacheng():
    print("Step 1: 段落修改（摘要前缀）")
    src = find_docx('029822410548')
    doc = Document(str(src))
    
    n = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        if text.startswith('摘要：') and text != '摘要：':
            remaining = 3
            for run in p.runs:
                if not run.text or remaining <= 0:
                    continue
                if len(run.text) <= remaining:
                    remaining -= len(run.text)
                    run.text = ''
                else:
                    run.text = run.text[remaining:]
                    remaining = 0
            n += 1
    print(f"  ✅ 删摘要前缀 {n}处")
    
    tmp = OUTPUT_DIR / "_tmp_fengjiacheng.docx"
    doc.save(str(tmp))
    print(f"  临时保存: {tmp}")
    return tmp


# ==================== 第二步：lxml修改XML结构 ====================

def step2_lizewen(tmp_path):
    print("\nStep 2: XML结构修改（tblGrid + Markdown表格替换）")
    
    # 用zipfile直接操作XML
    import zipfile, shutil
    
    final = OUTPUT_DIR / "029821410170_李泽文_修改版.docx"
    shutil.copy2(str(tmp_path), str(final))
    
    # 读取document.xml
    with zipfile.ZipFile(str(final), 'r') as z:
        doc_xml = etree.fromstring(z.read('word/document.xml'))
        other_files = {n: z.read(n) for n in z.namelist() if n != 'word/document.xml'}
    
    body_el = doc_xml.find(f'{{{NS}}}body')
    children = list(body_el)
    
    # 1. 修复tblGrid（找到第1个tbl）
    for child in children:
        tag = child.tag.split('}')[-1]
        if tag == 'tbl':
            tblGrid = child.find(f'{{{NS}}}tblGrid')
            if tblGrid is not None:
                gcs = list(tblGrid.findall(f'{{{NS}}}gridCol'))
                if len(gcs) == 1:
                    # 删旧的
                    for gc in gcs:
                        tblGrid.remove(gc)
                    # 加4列
                    for _ in range(4):
                        gc = etree.SubElement(tblGrid, f'{{{NS}}}gridCol')
                        gc.set(f'{{{NS}}}w', '1250')
                    # 修复每个tr的tc数
                    for tr in child.findall(f'{{{NS}}}tr'):
                        tcs = list(tr.findall(f'{{{NS}}}tc'))
                        while len(tcs) > 4:
                            tr.remove(tcs[-1])
                            tcs = list(tr.findall(f'{{{NS}}}tc'))
                    print("  ✅ 3.2节表格tblGrid修复")
            break
    
    # 2. 找到Markdown表格行并替换
    paras_to_remove = []
    for child in children:
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            texts = [t.text or '' for t in child.iter(f'{{{NS}}}t')
                    if t.tag == f'{{{NS}}}t']
            text = ''.join(texts).strip()
            if text.startswith('|') and any(kw in text for kw in ['问题类型', '---', '主题文化', '环保与舒适', '空间功能']):
                paras_to_remove.append(child)
    
    if paras_to_remove:
        first = paras_to_remove[0]
        idx = list(body_el).index(first)
        
        # 表标题
        title_p = etree.Element(f'{{{NS}}}p')
        pPr = etree.SubElement(title_p, f'{{{NS}}}pPr')
        jc = etree.SubElement(pPr, f'{{{NS}}}jc')
        jc.set(f'{{{NS}}}val', 'center')
        r = etree.SubElement(title_p, f'{{{NS}}}r')
        rPr = etree.SubElement(r, f'{{{NS}}}rPr')
        rf = etree.SubElement(rPr, f'{{{NS}}}rFonts')
        rf.set(f'{{{NS}}}ascii', '宋体')
        rf.set(f'{{{NS}}}eastAsia', '宋体')
        sz = etree.SubElement(rPr, f'{{{NS}}}sz')
        sz.set(f'{{{NS}}}val', '20')
        etree.SubElement(rPr, f'{{{NS}}}b')
        t = etree.SubElement(r, f'{{{NS}}}t')
        t.text = '表4.1 问题类型、实例与改进方向对应关系'
        body_el.insert(idx, title_p)
        
        # 表格
        new_tbl = make_table(
            ['问题类型', '典型实例', '改进方向（非解决方案，仅梳理思路）'],
            [
                ['主题文化表达流于表面', '某"农耕主题"民宿：仅放置旧农具，空间缺乏整体叙事', '强化空间叙事连贯性，材质与主题统一'],
                ['环保与舒适度冲突', '西南夯土民宿：未做内保温，湿度超75%，被褥发霉', '平衡传统外观与现代保温技术'],
                ['空间功能与需求错位', '观景书吧：书籍少、座椅硬，使用率不足30%', '基于游客行为模式重新规划功能'],
            ]
        )
        body_el.insert(idx, new_tbl)
        
        # 删Markdown行
        for elem in paras_to_remove:
            body_el.remove(elem)
        
        print(f"  ✅ Markdown表格→docx表格（表4.1，删{len(paras_to_remove)}行）")
    else:
        print("  ⚠️ 找不到Markdown表格")
    
    # 写回
    with zipfile.ZipFile(str(final), 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('word/document.xml', 
            etree.tostring(doc_xml, xml_declaration=True, encoding='UTF-8', standalone=True))
        for name, data in other_files.items():
            z.writestr(name, data)
    
    # 删临时
    os.remove(str(tmp_path))
    print(f"  📄 {final}")


def step2_fengjiacheng(tmp_path):
    print("\nStep 2: XML结构修改（4.5节表1替换）")
    
    import zipfile, shutil
    
    final = OUTPUT_DIR / "029822410548_冯嘉诚_修改版.docx"
    shutil.copy2(str(tmp_path), str(final))
    
    with zipfile.ZipFile(str(final), 'r') as z:
        doc_xml = etree.fromstring(z.read('word/document.xml'))
        other_files = {n: z.read(n) for n in z.namelist() if n != 'word/document.xml'}
    
    body_el = doc_xml.find(f'{{{NS}}}body')
    
    # 找到13行2列表格
    target = None
    for child in list(body_el):
        tag = child.tag.split('}')[-1]
        if tag == 'tbl':
            trs = child.findall(f'{{{NS}}}tr')
            if len(trs) >= 10:
                first_tr_tcs = list(trs[0].findall(f'{{{NS}}}tc'))
                if len(first_tr_tcs) == 2:
                    target = child
                    break
    
    if target is not None:
        new_tbl = make_table(
            ['智能化系统', '功能组成', '技术参数', '安装区域', '设计要点'],
            [
                ['紧急呼叫系统', '一键报警按钮\n拉绳报警器', '响应时间<5s\n拉绳高度300mm', '卫生间、卧室\n走廊', '双冗余设计\n声光联动报警'],
                ['环境监测系统', '温湿度传感器\nCO₂传感器\nPM2.5传感器', 'CO₂≤1000ppm\nPM2.5≤35μg/m³', '活动区、休息区\n公共走廊', '数据实时显示\n超标自动告警'],
                ['辅助照明系统', '智能筒灯\n人体感应灯\n低位夜间灯', '照度50-300lx\n色温3000K', '走廊、卫生间\n出入口', '红外+雷达双感\n缓亮缓灭'],
                ['智能控制系统', '中央控制面板\n分区分时控制', '支持远程/本地\n预设场景模式', '值班室\n公共区域', '一键场景切换\n能耗监测'],
            ]
        )
        idx = list(body_el).index(target)
        body_el.remove(target)
        body_el.insert(idx, new_tbl)
        print("  ✅ 4.5节表1已按智能化系统四要件重新组织")
    else:
        print("  ⚠️ 找不到目标表格")
    
    with zipfile.ZipFile(str(final), 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('word/document.xml',
            etree.tostring(doc_xml, xml_declaration=True, encoding='UTF-8', standalone=True))
        for name, data in other_files.items():
            z.writestr(name, data)
    
    os.remove(str(tmp_path))
    print(f"  📄 {final}")


if __name__ == '__main__':
    print("="*50)
    print("李泽文 (029821410170)")
    print("="*50)
    tmp = step1_lizewen()
    step2_lizewen(tmp)
    
    print(f"\n{'='*50}")
    print("冯嘉诚 (029822410548)")
    print("="*50)
    tmp = step1_fengjiacheng()
    step2_fengjiacheng(tmp)
