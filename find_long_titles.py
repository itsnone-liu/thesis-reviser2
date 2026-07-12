#!/usr/bin/env python3
"""找出状态表中论文名称过长的记录。"""
import openpyxl

BASE = "/root/project/workspace/论文资料/整理"
wb = openpyxl.load_workbook(f'{BASE}/学位申请合并_含论文状态.xlsx')
ws = wb['学位申请汇总']

print("论文名称过长的记录（>80字符）：")
count = 0
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, 1).value
    paper = ws.cell(r, 6).value
    major = ws.cell(r, 3).value
    data = ws.cell(r, 5).value
    
    if paper and len(str(paper)) > 80:
        count += 1
        sid = ws.cell(r, 2).value
        print(f'\n{count}. {name}({major}) 准考证={sid}')
        print(f'   当前名称(前80字): {str(paper)[:80]}...')
        print(f'   长度: {len(str(paper))}字符')

print(f'\n共 {count} 篇')
