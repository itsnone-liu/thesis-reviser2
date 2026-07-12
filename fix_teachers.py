#!/usr/bin/env python3
"""
指导教师刘焱冬/蒯瑶阳 → 替换为该专业其他老师。
改封面 + 更新状态表。
"""
import os, re
from docx import Document
import openpyxl

BASE = "/root/project/workspace/论文资料/整理"
XLSX = f"{BASE}/学位申请合并_含论文状态.xlsx"

# === 分配方案（专业 → 新老师） ===
NEW_TEACHERS = {
    "人力资源管理": "王娟",
    "财务管理": "张铭梅",
    "金融学": "张铭梅",
    # 环境设计有数据的只有张陆平，保留刘焱冬不动（3篇中2篇无文件）
}

# === 获取需要改的学生列表 ===
wb = openpyxl.load_workbook(XLSX)
ws = wb['学位申请汇总']

targets = []
for r in range(2, ws.max_row + 1):
    teacher = str(ws.cell(r, 7).value or '')
    if '刘焱冬' not in teacher and '蒯瑶阳' not in teacher:
        continue
    name = ws.cell(r, 1).value
    sid = str(int(ws.cell(r, 2).value))
    major = str(ws.cell(r, 3).value or '')
    data = str(ws.cell(r, 5).value or '')
    
    if data != '有':
        print(f"  ⏭ {name}({major}): 无文件，跳过")
        continue
    
    new_teacher = NEW_TEACHERS.get(major)
    if not new_teacher:
        print(f"  ⏭ {name}({major}): 无分配方案，跳过")
        continue
    
    targets.append((name, sid, major, new_teacher))

print(f"需修改的学生（有文件）: {len(targets)} 人\n")

# === 查找文件并修改封面 ===
def find_file(sid):
    for batch in ["25", "26"]:
        for major in ["经管", "机械", "设计"]:
            d = os.path.join(BASE, "论文终版", batch, major)
            if not os.path.exists(d):
                continue
            for f in os.listdir(d):
                if sid in f and f.endswith('.docx'):
                    return os.path.join(d, f)
    return None

for name, sid, major, new_teacher in targets:
    fp = find_file(sid)
    if not fp:
        print(f"  ⚠ {name}: 在论文终版中未找到文件")
        continue
    
    doc = Document(fp)
    modified = False
    
    for p in doc.paragraphs:
        text = p.text.strip()
        if '指导教师' in text or '指导老师' in text:
            for r in p.runs:
                if r.text and ('刘焱冬' in r.text or '蒯瑶阳' in r.text or 
                               '指导教师' in r.text or '指导老师' in r.text):
                    r.text = f"指导教师：{new_teacher}"
                    modified = True
                else:
                    r.text = ''
    
    if modified:
        doc.save(fp)
        print(f"  ✅ {name}({major}): 封面已修改 刘焱冬/蒯瑶阳 → {new_teacher}")
    else:
        print(f"  ⚠ {name}: 未找到教师字段")
    
    # 更新状态表
    for r in range(2, ws.max_row + 1):
        sid_raw = ws.cell(r, 2).value
        if sid_raw and str(int(sid_raw)) == sid:
            ws.cell(r, 7).value = new_teacher
            print(f"     → 状态表已更新")
            break

# === 保存状态表 ===
wb.save(XLSX)
print(f"\n✅ 全部完成！共处理 {len(targets)} 篇")
