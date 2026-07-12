#!/usr/bin/env python3
"""
环艺设计类论文批量修改脚本 v2
处理35篇低/无等级论文，针对审核表逐一修复

v2改进：
- 修复表格解析：适配设计类所有表格格式（<tr<th在同一行、无>标签等）
- 直接操作DOCX段落run文字保持格式
- 表格重建使用lxml操作XML层
"""

import os
import re
import sys
from pathlib import Path
from lxml import etree
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import openpyxl

# ==================== 路径配置 ====================
HUANYI_DIR = Path("/root/project/workspace/论文资料/整理/环艺")
MODIFIED_DIR = Path("/root/project/workspace/论文资料/整理/设计类_修改版")
OUTPUT_DIR = Path("/root/project/workspace/论文资料/整理/环艺_修改版")
AUDIT_FILE = Path("/root/project/workspace/论文资料/整理/设计类审核表.xlsx")

NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


# ==================== 辅助函数 ====================

def find_docx(exam_id):
    """找到环艺目录中的DOCX"""
    for f in os.listdir(HUANYI_DIR):
        if exam_id in f and f.endswith('.docx'):
            return HUANYI_DIR / f
    return None


def find_txt(exam_id):
    """找到修改版中的TXT"""
    for d in os.listdir(MODIFIED_DIR):
        if exam_id in d:
            txt = MODIFIED_DIR / d / "paper.txt"
            if txt.exists():
                return txt
    # 后6位模糊匹配
    short = exam_id[-6:]
    for d in os.listdir(MODIFIED_DIR):
        if short in d:
            txt = MODIFIED_DIR / d / "paper.txt"
            if txt.exists():
                return txt
    return None


def extract_cell_content(text):
    """从 <th内容</th 或 <td内容</td 提取纯内容"""
    for tag in ['<th', '<td', '<caption', '<tr', '</tr', '</th', '</td', '</table']:
        if text.startswith(tag):
            text = text[len(tag):]
        if text.endswith(tag):
            text = text[:-len(tag)]
    # 也去掉</tr和</t这种残缺尾巴
    for suffix in ['</tr', '</t', '</table']:
        if text.endswith(suffix):
            text = text[:-len(suffix)]
    return text.strip()


# ==================== TXT表格解析（核心） ====================

def parse_txt_tables(txt_path):
    """
    解析设计类TXT中的所有表格。
    支持多种格式：
    - <table /> + 换行 + <tr><td>（李泽文格式）
    - <table/>/> + <tr<thA</th<thB</th（毕守荣格式，同一行）
    - <table\n<tr<td...（无缩进格式）
    
    返回: [(标题, 表头列表, [数据行]), ...]
    """
    with open(txt_path, 'r') as f:
        lines = f.readlines()
    
    tables = []
    i = 0
    
    while i < len(lines):
        s = lines[i].strip()
        
        if not s.startswith('<table'):
            i += 1
            continue
        
        # --- 表格开始 ---
        rows_data = []  # [(is_header, [cells])]
        current_cells = []
        is_header = False
        caption = ''
        i += 1
        
        while i < len(lines):
            s = lines[i].strip()
            
            # 结束标签
            if s.startswith('</table'):
                if current_cells:
                    rows_data.append((is_header, current_cells))
                break
            
            # 行开始标签（可能在同一行有cell内容）
            if s.startswith('<tr'):
                # 保存上一行
                if current_cells:
                    rows_data.append((is_header, current_cells))
                
                current_cells = []
                is_header = False
                
                # 如果同一行有cell，提取它们
                # 按<th或<td分割
                parts = re.split(r'(?=<(?:th|td)\b)', s)
                # 但\b在中文前不匹配，试试去掉\b
                if len(parts) == 1:
                    parts = re.split(r'(?=<(?:th|td))', s)
                
                for part in parts:
                    if part.startswith('<th') or part.startswith('<td'):
                        is_header = is_header or part.startswith('<th')
                        content = extract_cell_content(part)
                        if content:
                            current_cells.append(content)
                        elif part.startswith('<td') and not part.startswith('<th'):
                            # <td可能只有标签没有内容
                            pass
                
                i += 1
                continue
            
            # 行结束标签
            if s.startswith('</tr') or s == '</tr':
                if current_cells:
                    rows_data.append((is_header, current_cells))
                current_cells = []
                is_header = False
                i += 1
                continue
            
            # 独立cell行（单独<th或<td在一行）
            if s.startswith('<th') or s.startswith('<td'):
                is_header = is_header or s.startswith('<th')
                content = extract_cell_content(s)
                current_cells.append(content)
                i += 1
                continue
            
            # 标题行
            if s.startswith('<caption'):
                # 提取标题
                m = re.search(r'<caption[^>]*>(.*?)</caption', s)
                if m:
                    caption = m.group(1).strip()
                else:
                    caption = s.replace('<caption', '').replace('</caption', '').strip()
                i += 1
                continue
            
            # 空白行或非表格内容
            if not s and current_cells:
                # 空行可能结束了一个未闭合的行
                pass
            
            i += 1
        
        # 处理parse结果，将行分为表头和数据行
        if not rows_data:
            i += 1
            continue
        
        headers = []
        data_rows = []
        
        for is_h, cells in rows_data:
            if is_h and not headers:
                headers = cells
            elif not is_h:
                if not headers:
                    headers = cells  # 第一行当表头
                elif cells:
                    data_rows.append(cells)
        
        # 验证：表头和数据行必须有数据
        if headers:
            tables.append((caption, headers, data_rows))
        
        i += 1
    
    return tables


# ==================== DOCX表格重建 ====================

def rebuild_table(table, headers, data_rows):
    """
    使用lxml直接修改表格的XML结构。
    保持table对象不变，只替换内容。
    """
    tbl = table._tbl
    
    existing_trs = list(tbl.findall(f'{{{NS}}}tr'))
    
    # 需要的总行数
    needed_rows = 1 + len(data_rows)
    num_cols = len(headers)
    
    # 删除多余行
    while len(existing_trs) > needed_rows:
        tbl.remove(existing_trs[-1])
        existing_trs = list(tbl.findall(f'{{{NS}}}tr'))
    
    # 增加缺少的行
    while len(existing_trs) < needed_rows:
        row_elem = etree.SubElement(tbl, f'{{{NS}}}tr')
        for _ in range(num_cols):
            tc = etree.SubElement(row_elem, f'{{{NS}}}tc')
            p = etree.SubElement(tc, f'{{{NS}}}p')
            r = etree.SubElement(p, f'{{{NS}}}r')
            t = etree.SubElement(r, f'{{{NS}}}t')
            t.text = ''
        existing_trs = list(tbl.findall(f'{{{NS}}}tr'))
    
    # 调整每行列数
    for ri, tr in enumerate(existing_trs):
        tcs = list(tr.findall(f'{{{NS}}}tc'))
        while len(tcs) > num_cols:
            tr.remove(tcs[-1])
            tcs = list(tr.findall(f'{{{NS}}}tc'))
        while len(tcs) < num_cols:
            tc = etree.SubElement(tr, f'{{{NS}}}tc')
            p = etree.SubElement(tc, f'{{{NS}}}p')
            r = etree.SubElement(p, f'{{{NS}}}r')
            t = etree.SubElement(r, f'{{{NS}}}t')
            t.text = ''
            tcs = list(tr.findall(f'{{{NS}}}tc'))
        
        # 填充数据
        for ci, tc in enumerate(tcs):
            # 清空段落
            for p in list(tc.findall(f'{{{NS}}}p')):
                tc.remove(p)
            
            # 新建段落
            p = etree.SubElement(tc, f'{{{NS}}}p')
            pPr = etree.SubElement(p, f'{{{NS}}}pPr')
            jc = etree.SubElement(pPr, f'{{{NS}}}jc')
            jc.set(f'{{{NS}}}val', 'center')
            
            # 获取该单元格的值
            if ri == 0 and ci < len(headers):
                val = headers[ci]
                bold = True
                size = 18  # 9pt
            elif ri > 0 and ri - 1 < len(data_rows) and ci < len(data_rows[ri-1]):
                val = data_rows[ri-1][ci]
                bold = False
                size = 18
            else:
                val = ''
                bold = False
                size = 18
            
            # 设置run格式
            r = etree.SubElement(p, f'{{{NS}}}r')
            rPr = etree.SubElement(r, f'{{{NS}}}rPr')
            rFonts = etree.SubElement(rPr, f'{{{NS}}}rFonts')
            rFonts.set(f'{{{NS}}}ascii', '宋体')
            rFonts.set(f'{{{NS}}}eastAsia', '宋体')
            sz = etree.SubElement(rPr, f'{{{NS}}}sz')
            sz.set(f'{{{NS}}}val', str(size))
            if bold:
                b = etree.SubElement(rPr, f'{{{NS}}}b')
            t = etree.SubElement(r, f'{{{NS}}}t')
            t.text = val if val else ''
    
    # 设置表格宽度100%
    tblPr = tbl.find(f'{{{NS}}}tblPr')
    if tblPr is None:
        tblPr = etree.Element(f'{{{NS}}}tblPr')
        tbl.insert(0, tblPr)
    tblW = tblPr.find(f'{{{NS}}}tblW')
    if tblW is None:
        tblW = etree.SubElement(tblPr, f'{{{NS}}}tblW')
    tblW.set(f'{{{NS}}}w', '5000')
    tblW.set(f'{{{NS}}}type', 'pct')
    
    # 删除表格边框（保持Table Grid风格不需要额外操作，docx默认有）
    
    return True


# ==================== 操作类型 ====================

def fix_abstract_prefix(doc):
    """删除摘要正文开头的'摘要：'或'摘要:'，保持run格式"""
    fixes = 0
    for p in doc.paragraphs:
        full_text = p.text.strip()
        if not (full_text.startswith('摘要：') or full_text.startswith('摘要:')):
            continue
        if full_text == '摘要：':
            continue
        
        prefix = '摘要：' if full_text.startswith('摘要：') else '摘要:'
        
        # 在run级别删除前缀
        remaining = len(prefix)
        for run in p.runs:
            if not run.text:
                continue
            if remaining <= 0:
                break
            if len(run.text) <= remaining:
                remaining -= len(run.text)
                run.text = ''
            else:
                run.text = run.text[remaining:]
                remaining = 0
        
        if remaining <= 0:
            fixes += 1
    
    return fixes


def fix_tables(doc, txt_path):
    """从TXT重新生成DOCX中的表格"""
    tables = parse_txt_tables(txt_path)
    if not tables:
        return "TXT中无表格"
    
    # 找到所有1列表格
    bad_indices = []
    for ti, table in enumerate(doc.tables):
        if len(table.columns) == 1 and len(table.rows) >= 3:
            bad_indices.append(ti)
    
    if not bad_indices:
        return f"无1列表格 (TXT有{len(tables)}个)"
    
    replaced = 0
    for bi in bad_indices:
        if replaced >= len(tables):
            break
        caption, headers, data_rows = tables[replaced]
        if not headers or len(headers) < 2:
            replaced += 1
            continue
        
        try:
            rebuild_table(doc.tables[bi], headers, data_rows)
            replaced += 1
        except Exception as e:
            return f"重建表格{bi}失败: {e}"
    
    return f"替换{replaced}/{len(bad_indices)}个表格 (TXT共{len(tables)}个)"


def fix_cover_level(doc):
    """修复层次：本科→专升本"""
    fixes = 0
    for p in doc.paragraphs:
        for run in p.runs:
            if '本科' in run.text and '层次' in ''.join(r.text for r in p.runs):
                run.text = run.text.replace('本科', '专升本')
                fixes += 1
    return fixes


def fix_cover_major(doc):
    """修复专业名称"""
    fixes = 0
    for p in doc.paragraphs:
        full = ''.join(r.text for r in p.runs)
        if '专' in full and '业' in full and '环境设计' not in full:
            for run in p.runs:
                m = re.search(r'(专\s*业\s*[：:]\s*)\S+', run.text)
                if m:
                    run.text = run.text.replace(m.group(1).strip(), f'专业：环境设计')
                    fixes += 1
    return fixes


def fix_cover_id(doc, correct_id):
    """修复封面学号"""
    fixes = 0
    for p in doc.paragraphs:
        full = ''.join(r.text for r in p.runs)
        if '学' in full and '号' in full:
            m = re.search(r'(\d{6,})', full)
            if m:
                orig_id = m.group(1)
                if orig_id != correct_id:
                    to_replace = orig_id[:len(correct_id)]
                    for run in p.runs:
                        if orig_id in run.text:
                            run.text = run.text.replace(orig_id, correct_id)
                            fixes += 1
    return fixes


def fix_text_replacement(doc, old, new):
    """批量替换文本（如 本研究团队→本研究）"""
    fixes = 0
    for p in doc.paragraphs:
        for run in p.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                fixes += 1
    return fixes


# ==================== 主处理流程 ====================

def process_paper(exam_id, name, problems):
    """处理单篇论文"""
    
    # 找文件
    src = find_docx(exam_id)
    if not src:
        return {'status': 'error', 'msg': f'找不到DOCX: {exam_id}'}
    
    txt_path = find_txt(exam_id)
    
    print(f"  处理: {src.name}")
    
    # 读取
    doc = Document(str(src))
    fixes = []
    
    # ====== 1. 删摘要前缀 ======
    if '摘要部分' in problems and ('多了' in problems or '摘要：' in problems):
        n = fix_abstract_prefix(doc)
        if n:
            fixes.append(f"删摘要前缀 {n}处")
    
    # ====== 2. 修复表格 ======
    if ('表格格式不对' in problems or '表格内容不对' in problems) and txt_path:
        result = fix_tables(doc, txt_path)
        fixes.append(f"表格: {result}")
    
    # ====== 3. 封面修正 ======
    if '层次' in problems and '专升本' in problems:
        n = fix_cover_level(doc)
        if n:
            fixes.append(f"层次→专升本 {n}处")
    
    if '专业名称不对' in problems or ('专业' in problems and '环境设计' in problems):
        n = fix_cover_major(doc)
        if n:
            fixes.append("专业→环境设计")
    
    if '学号' in problems:
        m = re.search(r'应为[：:"\u201c]*(\d{10,})', problems)
        correct = m.group(1) if m else exam_id
        n = fix_cover_id(doc, correct)
        if n:
            fixes.append(f"学号→{correct}")
    
    if '缺少学号' in problems:
        m = re.search(r'应为[：:"\u201c]*(\d{10,})', problems)
        if m:
            n = fix_cover_id(doc, m.group(1))
            if n:
                fixes.append(f"学号→{m.group(1)}")
    
    # ====== 4. 其他文本替换 ======
    if '本研究团队' in problems:
        n = fix_text_replacement(doc, '本研究团队', '本研究')
        if n:
            fixes.append(f"本研究团队→本研究 {n}处")
    
    # 目录页码修复（提示需要渲染）
    if '目录' in problems and '页码' in problems:
        fixes.append("目录页码: 需重新渲染（跳过）")
    
    # 正文格式（提示人工确认）
    if '正文格式不对' in problems or '正文格式不正确' in problems:
        fixes.append("正文格式: 需人工确认")
    
    # 保存
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_name = f"{exam_id}_{name}_修改版.docx"
    out_path = OUTPUT_DIR / out_name
    doc.save(str(out_path))
    
    return {'status': 'ok', 'fixes': fixes, 'output': str(out_path)}


def main():
    # 读取审核表
    wb = openpyxl.load_workbook(str(AUDIT_FILE))
    ws = wb['Sheet1']
    
    papers = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name_raw, exam_raw, _, problems, level = row
        exam_id = str(exam_raw or '').strip().zfill(12)
        lvl = str(level).strip() if level else 'None'
        if lvl not in ('低', 'None', 'None'):
            continue
        papers.append({
            'name': str(name_raw or '').strip(),
            'exam_id': exam_id,
            'problems': str(problems or '').strip(),
        })
    
    # 命令行筛选
    if len(sys.argv) > 1:
        papers = [p for p in papers if any(t in p['exam_id'] for t in sys.argv[1:])]
    
    print(f"需处理: {len(papers)}篇\n")
    
    results = []
    for i, p in enumerate(papers):
        name = p['name']
        eid = p['exam_id']
        problems = p['problems']
        
        print(f"[{i+1}/{len(papers)}] {name} ({eid})")
        for line in problems.split('\n'):
            if line.strip():
                print(f"  📋 {line.strip()[:80]}")
        
        r = process_paper(eid, name, problems)
        
        if r['status'] == 'error':
            print(f"  ❌ {r['msg']}")
        else:
            for f in r['fixes']:
                print(f"  ✅ {f}")
            print(f"  📄 {r['output']}")
        
        results.append(r)
        print()
    
    ok = sum(1 for r in results if r['status'] == 'ok')
    err = sum(1 for r in results if r['status'] == 'error')
    print(f"\n{'='*50}")
    print(f"完成！成功 {ok}, 失败 {err}")


if __name__ == '__main__':
    main()
