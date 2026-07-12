#!/usr/bin/env python3
"""
全面修复所有论文表格的序号和标题。
"""
import os, re, copy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

MODIFIED_DIR = "/root/project/workspace/论文资料/整理/环艺_修改版"
files = sorted([f for f in os.listdir(MODIFIED_DIR) if f.endswith('.docx')])

def add_title_before_table(doc, table, title_text):
    """在表格前插入标题段落。"""
    tbl_element = table._tbl
    parent = tbl_element.getparent()
    
    new_p = OxmlElement('w:p')
    
    # 段落属性
    pPr = OxmlElement('w:pPr')
    pStyle = OxmlElement('w:pStyle')
    pStyle.set(qn('w:val'), 'a7')  # 图表标题样式
    pPr.append(pStyle)
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    pPr.append(jc)
    new_p.append(pPr)
    
    # run
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    b = OxmlElement('w:b')
    rPr.append(b)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '24')
    rPr.append(sz)
    szCs = OxmlElement('w:szCs')
    szCs.set(qn('w:val'), '24')
    rPr.append(szCs)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = title_text
    r.append(t)
    new_p.append(r)
    
    parent.insert(list(parent).index(tbl_element), new_p)

def update_table_caption(doc, table, new_title):
    """替换表格前的标题文本。"""
    tbl_element = table._tbl
    prev_sib = tbl_element.getprevious()
    while prev_sib is not None:
        if prev_sib.tag == qn('w:p'):
            texts = [r.text or '' for r in prev_sib.findall('.//' + qn('w:t'))]
            full_text = ''.join(texts).strip()
            if re.search(r'表\s*\d', full_text[:30]):
                # 找到标题段落，清空所有文本再写入新标题
                for r in prev_sib.findall('.//' + qn('w:t')):
                    r.text = ''
                first_t = prev_sib.find('.//' + qn('w:t'))
                if first_t is not None:
                    first_t.text = new_title
                return True
        prev_sib = prev_sib.getprevious()
    return False

def check_table_issues(doc):
    """检查论文表格的序号问题。"""
    tables = doc.tables
    issues = []
    
    for t_idx, table in enumerate(tables):
        tbl_element = table._tbl
        prev_sib = tbl_element.getprevious()
        found_caption = False
        caption_num = None
        caption_text = ""
        
        while prev_sib is not None:
            if prev_sib.tag == qn('w:p'):
                texts = [r.text or '' for r in prev_sib.findall('.//' + qn('w:t'))]
                full_text = ''.join(texts).strip()
                m = re.search(r'(表\s*[\d.]+)(.*)', full_text[:50])
                if m:
                    found_caption = True
                    caption_num_str = m.group(1)
                    caption_text = m.group(2).strip() if m.group(2) else ""
                    # 提取数字部分
                    nums = re.findall(r'\d+', caption_num_str)
                    caption_num = '.'.join(nums)
                    break
            prev_sib = prev_sib.getprevious()
        
        issues.append({
            'idx': t_idx,
            'found': found_caption,
            'num': caption_num,
            'text': caption_text,
        })
    
    return issues

# =============================================
# 修复规则定义
# =============================================
repairs = {
    "029822410501_贾燕霞_修改版.docx": [
        ("add", 0, "表1 理论-设计要素映射表"),
        ("replace", 1, "表2 用户画像与设计响应"),
    ],
    "029823410134_高立凯_修改版.docx": [
        ("add", 0, "表1 理论与设计关联对照表"),
        ("replace", 1, "表2 材料与家具选型对比"),
    ],
    "029823410119_何素珍_修改版.docx": [
        ("add", 0, "表1 问题-策略映射表"),
        ("replace", 1, "表2 新型环保材料性能对比表"),
        ("replace", 2, "表3 主要材料清单"),
    ],
    "029823410164_臧斯恒_修改版.docx": [
        ("add", 0, "表1 理论-设计要素对照表"),
        ("add", 1, "表2 案例生态主题表现对比表"),
        ("replace", 2, "表3 材料参数与节能对比"),
    ],
    "029823410190_许秀敏_修改版.docx": [
        ("add", 0, "表1 国内外案例对比分析表"),
        ("replace", 1, "表2 管理创新模式对比"),
    ],
    "029823410338_毕守荣_修改版.docx": [
        ("add", 0, "表1 设计转译模型表"),
        ("replace", 1, "表2 乡土植物配置表"),
    ],
    "029823410272_王琛_修改版.docx": [
        ("add", 1, "表2 问题-需求对照表"),
        ("replace", 2, "表3 主要材料样品表"),
    ],
    "029823410558_张颖_修改版.docx": [
        ("add", 0, "表1 测点照度实测数据"),
        ("replace", 1, "表2 设计优化前后关键参数对比"),
        ("add", 2, "表3 设计成效对照表"),
    ],
    "029821410170_李泽文_修改版.docx": [
        # 表4.1 → 表1，表3.1 → 表2
        ("replace", 0, "表1 问题类型、实例与改进方向对应关系"),
        ("replace", 1, "表2 设计策略与问题对应表"),
    ],
}

# =============================================
# 执行修复
# =============================================
total_added = 0
total_replaced = 0
total_checked = 0

for fname in files:
    path = os.path.join(MODIFIED_DIR, fname)
    doc = Document(path)
    tables = doc.tables
    name = fname.split('_')[1]
    
    # 先检查是否有问题
    issues = check_table_issues(doc)
    
    # 检查多表格且序号不连续
    multi_table_issue = False
    if len(tables) > 1:
        # 检查是否所有表格都有标题且序号是1,2,3...
        all_titled_and_sequential = True
        for i, iss in enumerate(issues):
            if not iss['found']:
                all_titled_and_sequential = False
                break
            expected = str(i + 1)
            if iss['num'] != expected:
                all_titled_and_sequential = False
                break
        
        if not all_titled_and_sequential:
            multi_table_issue = True
    
    if fname in repairs:
        print(f"\n📄 {name}:")
        repair_rules = repairs[fname]
        modified = False
        
        for action, t_idx, title in repair_rules:
            if t_idx >= len(tables):
                print(f"  ⚠️  表格#{t_idx+1}不存在，跳过")
                continue
            
            if action == "add":
                add_title_before_table(doc, tables[t_idx], title)
                total_added += 1
                modified = True
                print(f"  ✅ 表格#{t_idx+1} → 添加 '{title}'")
            
            elif action == "replace":
                success = update_table_caption(doc, tables[t_idx], title)
                if success:
                    total_replaced += 1
                    modified = True
                    print(f"  ✅ 表格#{t_idx+1} → 更新为 '{title}'")
                else:
                    # 没找到标题，添加
                    add_title_before_table(doc, tables[t_idx], title)
                    total_added += 1
                    modified = True
                    print(f"  ✅ 表格#{t_idx+1} → 未找到原标题，添加 '{title}'")
        
        if modified:
            doc.save(path)
            print(f"  💾 已保存")
    
    total_checked += 1

print(f"\n{'='*50}")
print(f"📊 修复完成统计")
print(f"  ✅ 新增标题: {total_added}")
print(f"  ✅ 更新序号: {total_replaced}")
print(f"  📄 共检查 {total_checked} 篇论文")
print(f"  🔧 修改 {len(repairs)} 篇")
